#!/usr/bin/env python3
"""Shared TikHub Xiaohongshu search + Excel export for agent skills.

Secrets come only from environment variables injected by the skill runner.
Progress goes to stderr; callers print a single JSON object on stdout.
"""

from __future__ import annotations

import json
import os
import re
import sys
import threading
import time
from collections import Counter, OrderedDict, deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
)
SPLIT_RE = re.compile(r"[\n,，;；]+")
DEFAULT_PAGES = 10
DEFAULT_WORKERS = 5
DEFAULT_QPS = 10
MAX_PAGES = 20
MAX_WORKERS = 10

# App V2 这几条路由询价均为 $0.01/次，且不接受免费额度。
UNIT_USD = 0.01
DETAIL_IMAGE_PATH = "/api/v1/xiaohongshu/app_v2/get_image_note_detail"
DETAIL_VIDEO_PATH = "/api/v1/xiaohongshu/app_v2/get_video_note_detail"
USER_INFO_PATH = "/api/v1/xiaohongshu/app_v2/get_user_info"
SEARCH_NOTES_PATH = "/api/v1/xiaohongshu/app_v2/search_notes"
USER_POSTED_PATH = "/api/v1/xiaohongshu/app_v2/get_user_posted_notes"
MAX_POSTED_PAGES = 50
USER_PROFILE_RE = re.compile(r"xiaohongshu\.com/user/profile/([0-9a-fA-F]{24})", re.I)

URL_RE = re.compile(r"https?://[^\s\"'<>\\）)】\]]{8,}", re.I)
NOTE_ID_ONLY_RE = re.compile(r"^[0-9a-fA-F]{24}$")
NOTE_ID_IN_URL_RES = (
    re.compile(
        r"xiaohongshu\.com/(?:explore|discovery/item|item|notes?/item|note)/([0-9a-fA-F]{24})",
        re.I,
    ),
    re.compile(r"xiaohongshu\.com/user/profile/[0-9a-fA-F]+/([0-9a-fA-F]{24})", re.I),
    re.compile(r"[?&](?:note_id|noteId|noteid)=([0-9a-fA-F]{24})", re.I),
)


def is_fatal_tikhub(err: Exception | str) -> bool:
    text = str(err)
    return any(code in text for code in ("401", "402", "403"))


def uniq_warnings(items: list[str], limit: int = 6) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for item in items:
        text = (item or "").strip()
        if not text or text in seen:
            continue
        seen.add(text)
        out.append(text)
        if len(out) >= limit:
            break
    return out


class XhsError(Exception):
    """User-facing Xiaohongshu export failure."""


class RateLimiter:
    def __init__(self, qps: int) -> None:
        self.qps = max(1, int(qps))
        self._lock = threading.Lock()
        self._hits: deque[float] = deque()

    def acquire(self) -> None:
        while True:
            with self._lock:
                now = time.monotonic()
                while self._hits and now - self._hits[0] >= 1.0:
                    self._hits.popleft()
                if len(self._hits) < self.qps:
                    self._hits.append(now)
                    return
                wait = 1.0 - (now - self._hits[0])
            time.sleep(max(wait, 0.02))


def require_env(name: str) -> str:
    value = (os.environ.get(name) or "").strip()
    if not value:
        raise XhsError(f"missing env: {name}")
    return value


def optional_env(name: str) -> str:
    return (os.environ.get(name) or "").strip()


def proxy_host(proxy_url: str) -> str:
    rest = proxy_url.split("://", 1)[-1]
    return rest.split("@")[-1].split("/")[0] or "on"


def tikhub_error_text(resp: httpx.Response) -> str:
    try:
        data = resp.json()
    except Exception:  # noqa: BLE001
        return f"HTTP {resp.status_code}"
    detail = data.get("detail") if isinstance(data, dict) and isinstance(data.get("detail"), dict) else {}
    msg = (
        (detail.get("message_zh") if isinstance(detail, dict) else None)
        or (detail.get("message") if isinstance(detail, dict) else None)
        or (data.get("message_zh") if isinstance(data, dict) else None)
        or (data.get("message") if isinstance(data, dict) else None)
        or (data.get("msg") if isinstance(data, dict) else None)
        or resp.text[:200]
    )
    code = (detail.get("code") if isinstance(detail, dict) else None) or (
        data.get("code") if isinstance(data, dict) else None
    ) or resp.status_code
    return f"TikHub {code}: {msg}"


def parse_keywords(raw: Any) -> list[str]:
    parts: list[str] = []
    if isinstance(raw, list):
        for item in raw:
            parts.extend(SPLIT_RE.split(str(item or "")))
    elif isinstance(raw, str):
        parts.extend(SPLIT_RE.split(raw))
    out: list[str] = []
    seen: set[str] = set()
    for part in parts:
        kw = part.strip()
        if kw and kw not in seen:
            seen.add(kw)
            out.append(kw)
    return out


def num(v: Any) -> int:
    if v is None or v == "":
        return 0
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip().replace(",", "")
    try:
        if s.endswith("万"):
            return int(float(s[:-1]) * 10000)
        if s.endswith("亿"):
            return int(float(s[:-1]) * 100000000)
        return int(float(s))
    except Exception:
        return 0


def first_present_num(*values: Any) -> int:
    for value in values:
        if value is None or value == "":
            continue
        return num(value)
    return 0


def ensure_xlsx_filename(name: str) -> str:
    stem = Path(str(name or "").strip()).name
    if not stem:
        return stem
    lower = stem.lower()
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        return stem
    if lower.endswith(".csv") or lower.endswith(".xls"):
        return f"{stem.rsplit('.', 1)[0]}.xlsx"
    return f"{stem}.xlsx"


def fmt_time(ts: Any) -> str:
    try:
        ts = int(ts)
        if ts > 10**12:
            ts //= 1000
        return datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return ""


def media_type(t: str) -> str:
    t = (t or "").lower()
    if t in ("video", "video_note", "短视频"):
        return "短视频"
    if "live" in t or t == "直播笔记":
        return "直播"
    return "图文"


def detail_path(note_type: str) -> str:
    if (note_type or "").lower() == "video":
        return DETAIL_VIDEO_PATH
    return DETAIL_IMAGE_PATH


def extract_urls(text: str) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for match in URL_RE.findall(text or ""):
        url = match.rstrip(".,;，。；")
        if url and url not in seen:
            seen.add(url)
            out.append(url)
    return out


def extract_note_id_from_text(text: str) -> str:
    raw = (text or "").strip()
    if not raw:
        return ""
    if NOTE_ID_ONLY_RE.fullmatch(raw):
        return raw.lower()
    for pat in NOTE_ID_IN_URL_RES:
        match = pat.search(raw)
        if match:
            return match.group(1).lower()
    match = re.search(r"xiaohongshu\.com[^0-9a-fA-F]{0,24}([0-9a-fA-F]{24})", raw, re.I)
    if match:
        return match.group(1).lower()
    return ""


def guess_note_kind(*parts: str) -> str:
    blob = " ".join(str(p or "") for p in parts).lower()
    if any(token in blob for token in ("/video/", "video_note", "type=video", "短视频")):
        return "video"
    if re.search(r"(?:^|[\s,，/])视频(?:$|[\s,，])", blob) or blob.strip() in {"视频", "video"}:
        return "video"
    return "image"


def extract_detail_metrics(data: dict[str, Any]) -> dict[str, Any]:
    found = walk_first(
        data,
        {
            "note_id",
            "title",
            "display_title",
            "desc",
            "type",
            "liked_count",
            "nice_count",
            "collected_count",
            "comments_count",
            "comment_count",
            "shared_count",
            "share_count",
            "nickname",
            "timestamp",
            "time",
            "userid",
            "user_id",
        },
    )
    title = found.get("title")
    display = found.get("display_title")
    desc = found.get("desc")
    return {
        "note_id": str(found.get("note_id") or "").strip(),
        "type": found.get("type") or "",
        "title": (title.strip() if isinstance(title, str) else "")
        or (display.strip() if isinstance(display, str) else ""),
        "desc": desc.strip() if isinstance(desc, str) else "",
        "nickname": str(found.get("nickname") or "").strip(),
        "user_id": str(found.get("userid") or found.get("user_id") or "").strip(),
        "liked": first_present_num(found.get("liked_count"), found.get("likes"), found.get("nice_count")),
        "comments": first_present_num(found.get("comments_count"), found.get("comment_count")),
        "collects": num(found.get("collected_count")),
        "shares": first_present_num(found.get("shared_count"), found.get("share_count")),
        "time": fmt_time(found.get("timestamp") or found.get("time")),
        "body": "",
    }


def walk_first(obj: Any, names: set[str]) -> dict[str, Any]:
    found: dict[str, Any] = {}
    stack = [obj]
    while stack and len(found) < len(names):
        x = stack.pop()
        if isinstance(x, dict):
            for k, v in x.items():
                if k in names and k not in found:
                    found[k] = v
            stack.extend(x.values())
        elif isinstance(x, list):
            stack.extend(x[:4])
    return found


class TikHub:
    def __init__(self, key: str, base_url: str, qps: int, proxy: str = "") -> None:
        self.key = key
        self.base = base_url.rstrip("/")
        self.limiter = RateLimiter(qps)
        self.billed: Counter[str] = Counter()
        kwargs: dict[str, Any] = {
            "timeout": 45.0,
            "headers": {"Authorization": f"Bearer {key}", "User-Agent": UA},
        }
        if proxy:
            kwargs["proxy"] = proxy
        self.client = httpx.Client(**kwargs)

    def close(self) -> None:
        self.client.close()

    def get(self, path: str, params: dict[str, Any]) -> dict[str, Any]:
        last: Exception | None = None
        for attempt in range(4):
            self.limiter.acquire()
            try:
                resp = self.client.get(f"{self.base}{path}", params=params)
                if resp.status_code in (401, 402, 403):
                    raise XhsError(tikhub_error_text(resp))
                if resp.status_code in (429, 500, 502, 503, 504):
                    last = XhsError(tikhub_error_text(resp))
                    time.sleep(1.2 * (attempt + 1))
                    continue
                if resp.status_code >= 400:
                    raise XhsError(tikhub_error_text(resp))
                data = resp.json()
                if not isinstance(data, dict):
                    raise XhsError("TikHub 返回非对象")
                self.billed[path] += 1
                return data
            except XhsError:
                raise
            except Exception as err:  # noqa: BLE001
                last = err
                time.sleep(0.8 * (attempt + 1))
        raise XhsError(f"TikHub 请求失败: {last}")


def extract_note(item: dict[str, Any], keyword: str) -> dict[str, Any] | None:
    n = item.get("note") if isinstance(item, dict) else None
    if not isinstance(n, dict):
        n = item if isinstance(item, dict) else {}
    user = n.get("user") if isinstance(n.get("user"), dict) else {}
    nid = str(n.get("id") or n.get("note_id") or "").strip()
    if not nid:
        return None
    return {
        "note_id": nid,
        "type": n.get("type") or "",
        "title": (n.get("title") or n.get("display_title") or "").strip(),
        "desc": (n.get("desc") or "").strip(),
        "body": "",
        "nickname": user.get("nickname") or user.get("name") or "",
        "user_id": str(user.get("userid") or user.get("user_id") or user.get("id") or ""),
        "liked": first_present_num(n.get("liked_count"), n.get("likes"), n.get("nice_count")),
        "comments": first_present_num(n.get("comments_count"), n.get("comment_count")),
        "collects": num(n.get("collected_count")),
        "shares": first_present_num(n.get("shared_count"), n.get("share_count")),
        "time": fmt_time(n.get("timestamp") or n.get("time") or n.get("update_time")),
        "keywords": [keyword],
    }


def search_keyword(api: TikHub, keyword: str, pages: int) -> tuple[list[dict[str, Any]], str | None]:
    rows: list[dict[str, Any]] = []
    search_id = ""
    search_session_id = ""
    for page in range(1, pages + 1):
        params: dict[str, Any] = {
            "keyword": keyword,
            "page": page,
            "sort_type": "general",
            "note_type": "不限",
            "time_filter": "不限",
            "source": "explore_feed",
            "ai_mode": 0,
        }
        if search_id:
            params["search_id"] = search_id
        if search_session_id:
            params["search_session_id"] = search_session_id
        try:
            data = api.get("/api/v1/xiaohongshu/app_v2/search_notes", params)
        except XhsError as err:
            print(f"SEARCH_STOP keyword={keyword} page={page} {err}", file=sys.stderr, flush=True)
            return rows, str(err)
        if data.get("code") != 200:
            msg = data.get("message_zh") or data.get("message") or data.get("msg") or ""
            err = f"TikHub {data.get('code')}: {msg}".strip()
            print(f"SEARCH_STOP keyword={keyword} page={page} {err}", file=sys.stderr, flush=True)
            return rows, err
        inner = data.get("data") or {}
        search_id = inner.get("search_id") or search_id
        search_session_id = inner.get("search_session_id") or search_session_id
        payload = inner.get("data") if isinstance(inner.get("data"), dict) else {}
        items = payload.get("items") or []
        added = 0
        for it in items:
            rec = extract_note(it, keyword)
            if rec:
                rows.append(rec)
                added += 1
        print(f"page keyword={keyword} p={page} got={len(items)} keep={added}", file=sys.stderr, flush=True)
        if not items:
            break
    return rows, None


def merge_rows(batches: list[list[dict[str, Any]]]) -> OrderedDict[str, dict[str, Any]]:
    merged: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for batch in batches:
        for rec in batch:
            nid = rec["note_id"]
            if nid in merged:
                for kw in rec["keywords"]:
                    if kw not in merged[nid]["keywords"]:
                        merged[nid]["keywords"].append(kw)
                continue
            merged[nid] = rec
    return merged


def apply_detail(rec: dict[str, Any], data: dict[str, Any], fill_body: bool) -> None:
    found = walk_first(data, {"title", "desc"})
    title = found.get("title")
    desc = found.get("desc")
    if isinstance(title, str) and title.strip() and not rec["title"]:
        rec["title"] = title.strip()
    if isinstance(desc, str) and desc.strip():
        text = desc.strip()
        if fill_body:
            rec["body"] = text
        if not rec["desc"] or len(text) > len(rec["desc"]):
            rec["desc"] = text


def fill_details(
    api: TikHub, rows: OrderedDict[str, dict[str, Any]], with_body: bool, workers: int
) -> tuple[int, int, list[str]]:
    targets = [r for r in rows.values() if with_body or not r["title"]]
    ok = 0
    fail = 0
    warnings: list[str] = []

    def one(rec: dict[str, Any]) -> str:
        data = api.get(detail_path(str(rec["type"])), {"note_id": rec["note_id"]})
        if data.get("code") != 200:
            return "fail"
        apply_detail(rec, data, fill_body=with_body)
        if with_body and not rec["body"]:
            rec["body"] = rec["desc"]
        return "ok"

    if not targets:
        return 0, 0, []
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futs = {pool.submit(one, rec): rec["note_id"] for rec in targets}
        done = 0
        for fut in as_completed(futs):
            done += 1
            try:
                status = fut.result()
            except Exception as err:  # noqa: BLE001
                print(f"DETAIL_FAIL {futs[fut]} {err}", file=sys.stderr, flush=True)
                fail += 1
                warnings.append(str(err))
                continue
            if status == "ok":
                ok += 1
            else:
                fail += 1
            if done % 20 == 0 or done == len(targets):
                print(f"detail {done}/{len(targets)} ok={ok} fail={fail}", file=sys.stderr, flush=True)
    return ok, fail, warnings


def fill_fans(
    api: TikHub, rows: OrderedDict[str, dict[str, Any]], workers: int
) -> tuple[dict[str, int], list[str]]:
    uids: list[str] = []
    seen: set[str] = set()
    for rec in rows.values():
        uid = rec.get("user_id") or ""
        if uid and uid not in seen:
            seen.add(uid)
            uids.append(uid)
    fans_map: dict[str, int] = {uid: 0 for uid in uids}
    ok = 0
    warnings: list[str] = []

    def one(uid: str) -> tuple[str, int]:
        data = api.get("/api/v1/xiaohongshu/app_v2/get_user_info", {"user_id": uid})
        found = walk_first(data, {"fans"})
        return uid, num(found.get("fans"))

    if not uids:
        return fans_map, []
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futs = {pool.submit(one, uid): uid for uid in uids}
        done = 0
        for fut in as_completed(futs):
            done += 1
            uid = futs[fut]
            try:
                _, fans = fut.result()
                fans_map[uid] = fans
                if fans:
                    ok += 1
            except Exception as err:  # noqa: BLE001
                print(f"USER_FAIL {uid} {err}", file=sys.stderr, flush=True)
                warnings.append(str(err))
            if done % 20 == 0 or done == len(uids):
                print(f"users {done}/{len(uids)} fans_ok={ok}", file=sys.stderr, flush=True)
    return fans_map, warnings


def excel_headers(with_body: bool, with_fans: bool) -> list[str]:
    headers = ["时间", "来源", "作者号", "媒体类型", "标题", "地址", "摘要"]
    if with_body:
        headers.append("正文")
    if with_fans:
        headers.append("粉丝数")
    headers.extend(["转发数", "评论数", "点赞数", "收藏数", "搜索关键词"])
    return headers


def _cost_payload(rows: list[dict[str, Any]]) -> dict[str, Any]:
    total = round(sum(float(row["subtotal_usd"]) for row in rows), 4)
    lines = [
        "| 调用 | 次数 | 单价 | 小计 |",
        "|---|---:|---:|---:|",
    ]
    for row in rows:
        lines.append(
            f"| {row['name']} | {row['count']} | ${row['unit_usd']:.2f} | **${row['subtotal_usd']:.2f}** |"
        )
    if not rows:
        lines.append("| （无已计费请求） | 0 | $0.01 | **$0.00** |")
    lines.append(f"| **合计** | | | **${total:.2f}** |")
    return {
        "unit_usd": UNIT_USD,
        "currency": "USD",
        "calls": rows,
        "total_usd": total,
        "table": "\n".join(lines),
        "note": "按 HTTP 200 计费，402/失败不扣费。小红书 App V2 单价 $0.01/次。",
    }


def build_cost_from_paths(billed: Counter[str], labels: list[tuple[str, str]]) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    for path, label in labels:
        n = int(billed.get(path) or 0)
        if n <= 0:
            continue
        rows.append(
            {
                "name": label,
                "count": n,
                "unit_usd": UNIT_USD,
                "subtotal_usd": round(n * UNIT_USD, 4),
            }
        )
    return _cost_payload(rows)


def build_cost(billed: Counter[str], with_body: bool) -> dict[str, Any]:
    detail_count = billed.get(DETAIL_IMAGE_PATH, 0) + billed.get(DETAIL_VIDEO_PATH, 0)
    synthetic: Counter[str] = Counter(
        {
            SEARCH_NOTES_PATH: billed.get(SEARCH_NOTES_PATH, 0),
            "__detail__": detail_count,
            USER_INFO_PATH: billed.get(USER_INFO_PATH, 0),
        }
    )
    return build_cost_from_paths(
        synthetic,
        [
            (SEARCH_NOTES_PATH, "search_notes 翻页"),
            ("__detail__", "补正文（详情）" if with_body else "补空标题（详情）"),
            (USER_INFO_PATH, "get_user_info 拉粉丝"),
        ],
    )


def write_excel(
    path: Path,
    rows: OrderedDict[str, dict[str, Any]],
    fans_map: dict[str, int],
    with_body: bool,
    with_fans: bool,
) -> None:
    headers = excel_headers(with_body, with_fans)
    wb = Workbook()
    ws = wb.active
    ws.title = "文章列表"
    ws.append(headers)
    head_font = Font(bold=True)
    fill = PatternFill("solid", fgColor="F2F2F2")
    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.font = head_font
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")

    for rec in rows.values():
        url = f"https://www.xiaohongshu.com/discovery/item/{rec['note_id']}"
        line: list[Any] = [
            rec["time"],
            "小红书",
            rec["nickname"],
            media_type(str(rec["type"])),
            rec["title"] or (rec["desc"][:40] if rec["desc"] else rec["note_id"]),
            url,
            rec["desc"],
        ]
        if with_body:
            line.append(rec.get("body") or rec["desc"])
        if with_fans:
            line.append(fans_map.get(rec.get("user_id") or "", 0) or 0)
        line.extend(
            [
                rec["shares"],
                rec["comments"],
                rec["liked"],
                rec["collects"],
                " | ".join(rec["keywords"]),
            ]
        )
        ws.append(line)

    last = chr(ord("A") + len(headers) - 1)
    ws.auto_filter.ref = f"A1:{last}{ws.max_row}"
    ws.freeze_panes = "A2"
    widths = {
        "A": 20,
        "B": 10,
        "C": 18,
        "D": 10,
        "E": 42,
        "F": 55,
        "G": 50,
        "H": 50 if with_body else (10 if with_fans else 10),
        "I": 10,
        "J": 10,
        "K": 10,
        "L": 10,
        "M": 18,
        "N": 18,
    }
    for col, width in widths.items():
        if ord(col) - ord("A") + 1 <= len(headers):
            ws.column_dimensions[col].width = width
    wrap_cols = {7, 8} if with_body else {7}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column in wrap_cols)
            cell.border = thin
    ws.row_dimensions[1].height = 22
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def default_filename(keywords: list[str], with_body: bool, with_fans: bool) -> str:
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    slug = re.sub(r"[\\/:*?\"<>|]", "_", (keywords[0] if keywords else "search")[:20])
    extra = f"等{len(keywords)}词" if len(keywords) > 1 else ""
    kind = "带正文" if with_body else "无正文"
    if not with_fans:
        kind += "无粉丝"
    return f"小红书搜索-{kind}-{slug}{extra}-{stamp}.xlsx"


def run_export(params: dict[str, Any], with_body: bool, with_fans: bool = True) -> dict[str, Any]:
    keywords = parse_keywords(params.get("keywords") or params.get("keyword"))
    if not keywords:
        raise XhsError("必须提供 keywords（字符串或字符串数组）")
    try:
        pages = int(params.get("pages") or DEFAULT_PAGES)
    except (TypeError, ValueError) as err:
        raise XhsError("pages 必须是整数") from err
    if pages < 1 or pages > MAX_PAGES:
        raise XhsError(f"pages 必须在 1–{MAX_PAGES} 之间")
    try:
        workers = int(params.get("workers") or DEFAULT_WORKERS)
    except (TypeError, ValueError) as err:
        raise XhsError("workers 必须是整数") from err
    workers = max(1, min(workers, MAX_WORKERS, DEFAULT_QPS))
    export_dir = Path(params.get("export_dir") or "./workspace")
    filename = ensure_xlsx_filename(
        str(params.get("filename") or "").strip()
        or default_filename(keywords, with_body, with_fans)
    )
    if "/" in filename or "\\" in filename:
        raise XhsError("filename 只能是文件名，不能含路径")

    proxy = optional_env("SKILL_TIKHUB_PROXY")
    api = TikHub(
        require_env("SKILL_TIKHUB_API_KEY"),
        require_env("SKILL_TIKHUB_BASE_URL"),
        DEFAULT_QPS,
        proxy=proxy,
    )
    try:
        print(
            f"keywords={len(keywords)} pages={pages} workers={workers} "
            f"with_body={with_body} with_fans={with_fans}"
            + (f" proxy={proxy_host(proxy)}" if proxy else " proxy=off"),
            file=sys.stderr,
            flush=True,
        )
        batches: list[list[dict[str, Any]]] = []
        per_keyword: dict[str, int] = {}
        warnings: list[str] = []
        with ThreadPoolExecutor(max_workers=min(workers, len(keywords))) as pool:
            futs = {pool.submit(search_keyword, api, kw, pages): kw for kw in keywords}
            for fut in as_completed(futs):
                kw = futs[fut]
                try:
                    batch, err = fut.result()
                except Exception as err:  # noqa: BLE001
                    print(f"KEYWORD_FAIL {kw} {err}", file=sys.stderr, flush=True)
                    warnings.append(f"{kw}: {err}")
                    batch, err = [], str(err)
                per_keyword[kw] = len(batch)
                batches.append(batch)
                if err:
                    kept = f"，已保留前 {len(batch)} 条" if batch else ""
                    warnings.append(f"{kw}: {err}{kept}")
        merged = merge_rows(batches)
        raw_total = sum(len(batch) for batch in batches)
        print(f"DEDUP {raw_total} -> {len(merged)}", file=sys.stderr, flush=True)
        if raw_total == 0 and warnings:
            raise XhsError("；".join(uniq_warnings(warnings)))
        detail_ok, detail_fail, detail_warns = fill_details(
            api, merged, with_body=with_body, workers=workers
        )
        fans_map: dict[str, int] = {}
        fan_warns: list[str] = []
        if with_fans:
            fans_map, fan_warns = fill_fans(api, merged, workers=workers)
        warnings.extend(detail_warns)
        warnings.extend(fan_warns)
        if detail_fail:
            warnings.append(f"详情补全失败 {detail_fail} 条，对应列可能为空")
        if fan_warns:
            warnings.append("部分作者粉丝数未取到")
        out_path = (export_dir / filename).resolve()
        write_excel(out_path, merged, fans_map, with_body=with_body, with_fans=with_fans)
        cost = build_cost(api.billed, with_body=with_body)
    finally:
        api.close()

    types = Counter(media_type(str(rec["type"])) for rec in merged.values())
    warn_list = uniq_warnings(warnings)
    note = "请点击下载按钮获取 Excel。"
    if warn_list:
        note = (
            f"中途中断（常见原因：TikHub 付费余额不足），已导出已获取的 {len(merged)} 条。"
            f"{'；'.join(warn_list)} 请点击下载按钮获取 Excel。"
        )
    note += f" 本次消耗约 ${cost['total_usd']:.2f}，回复中须贴出 data.cost.table。"
    return {
        "output_xlsx": out_path.name,
        "output_path": str(out_path),
        "keywords": keywords,
        "pages": pages,
        "per_keyword": per_keyword,
        "raw_rows": raw_total,
        "rows": len(merged),
        "types": dict(types),
        "detail_ok": detail_ok,
        "detail_fail": detail_fail,
        "empty_title": sum(1 for rec in merged.values() if not rec["title"]),
        "empty_body": sum(1 for rec in merged.values() if with_body and not (rec.get("body") or rec.get("desc"))),
        "with_fans": with_fans,
        "partial": bool(warn_list),
        "warnings": warn_list,
        "cost": cost,
        "note": note,
    }


def emit_main(with_body: bool, with_fans: bool = True) -> None:
    raw = sys.argv[1] if len(sys.argv) > 1 else ""
    try:
        params = json.loads(raw) if str(raw).strip() else {}
        if not isinstance(params, dict):
            raise XhsError("参数必须是 JSON 对象")
        data = run_export(params, with_body=with_body, with_fans=with_fans)
        print(json.dumps({"ok": True, "data": data}, ensure_ascii=False))
    except XhsError as err:
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": str(err),
                    "hint": "提供 keywords；确认 SKILL_TIKHUB_API_KEY / SKILL_TIKHUB_BASE_URL；59 需 SKILL_TIKHUB_PROXY；402 表示该路由要付费余额",
                },
                ensure_ascii=False,
            )
        )
    except Exception as err:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": f"{type(err).__name__}: {err}"}, ensure_ascii=False))


def parse_user_targets(params: dict[str, Any]) -> list[dict[str, str]]:
    raw: list[str] = []
    for key in ("user_ids", "user_id", "account_id", "account_ids", "share_text"):
        value = params.get(key)
        if value is None or value == "":
            continue
        if isinstance(value, list):
            raw.extend(str(x) for x in value)
        else:
            raw.extend(SPLIT_RE.split(str(value)))
    targets: list[dict[str, str]] = []
    seen: set[str] = set()
    for part in raw:
        text = (part or "").strip()
        if not text:
            continue
        user_id = ""
        share = ""
        profile = USER_PROFILE_RE.search(text)
        if profile:
            user_id = profile.group(1).lower()
        elif NOTE_ID_ONLY_RE.fullmatch(text):
            user_id = text.lower()
        elif "xhslink." in text.lower() or "xiaohongshu.com" in text.lower():
            share = text
        else:
            user_id = text
        label = user_id or share
        if not label or label in seen:
            continue
        seen.add(label)
        targets.append({"label": label, "user_id": user_id, "share_text": share})
    return targets


def extract_posted_note(item: dict[str, Any], label: str) -> dict[str, Any] | None:
    n = item.get("note") if isinstance(item, dict) and isinstance(item.get("note"), dict) else item
    if not isinstance(n, dict):
        return None
    user = n.get("user") if isinstance(n.get("user"), dict) else {}
    nid = str(n.get("id") or n.get("note_id") or "").strip()
    if not nid:
        return None
    return {
        "note_id": nid,
        "type": n.get("type") or "",
        "title": (n.get("title") or n.get("display_title") or "").strip(),
        "desc": (n.get("desc") or "").strip(),
        "body": "",
        "nickname": user.get("nickname") or user.get("name") or "",
        "user_id": str(user.get("userid") or user.get("user_id") or user.get("id") or ""),
        "liked": first_present_num(n.get("liked_count"), n.get("likes"), n.get("nice_count")),
        "comments": first_present_num(n.get("comments_count"), n.get("comment_count")),
        "collects": num(n.get("collected_count")),
        "shares": first_present_num(n.get("shared_count"), n.get("share_count")),
        "time": fmt_time(
            n.get("create_time") or n.get("timestamp") or n.get("time") or n.get("last_update_time")
        ),
        "keywords": [label],
    }


def posted_upstream_error(data: dict[str, Any]) -> str | None:
    inner = data.get("data")
    blob_parts: list[str] = []
    if isinstance(inner, str):
        blob_parts.append(inner)
    elif isinstance(inner, dict):
        if inner.get("success") is False:
            blob_parts.append(str(inner.get("msg") or inner.get("message") or "服务异常"))
        for key in ("msg", "message", "message_zh"):
            val = inner.get(key)
            if isinstance(val, str) and val.strip():
                blob_parts.append(val)
        nested = inner.get("data")
        if isinstance(nested, str):
            blob_parts.append(nested)
    blob = " ".join(blob_parts)
    if any(token in blob for token in ("服务异常", "用户不存在", "账号不存在", "解析失败", "不存在该用户")):
        return blob.strip()[:200]
    if isinstance(inner, dict) and inner.get("success") is False:
        return (blob or "上游失败").strip()[:200]
    return None


def fetch_user_posted(
    api: TikHub, target: dict[str, str], pages: int, stop: threading.Event | None = None
) -> tuple[list[dict[str, Any]], str | None]:
    rows: list[dict[str, Any]] = []
    cursor = ""
    user_id = target.get("user_id") or ""
    share = "" if user_id else (target.get("share_text") or "")
    label = target.get("label") or user_id or share
    for page in range(1, pages + 1):
        if stop is not None and stop.is_set():
            return rows, "TikHub 中断（常见原因：付费余额不足）"
        params: dict[str, Any] = {}
        if user_id:
            params["user_id"] = user_id
        elif share:
            params["share_text"] = share
        else:
            return rows, "缺少 user_id 或 share_text"
        if cursor:
            params["cursor"] = cursor
        try:
            data = api.get(USER_POSTED_PATH, params)
        except XhsError as err:
            if is_fatal_tikhub(err) and stop is not None:
                stop.set()
            print(f"POSTED_STOP user={label} page={page} {err}", file=sys.stderr, flush=True)
            return rows, str(err)
        if data.get("code") != 200:
            msg = data.get("message_zh") or data.get("message") or data.get("msg") or ""
            err = f"TikHub {data.get('code')}: {msg}".strip()
            if is_fatal_tikhub(err) and stop is not None:
                stop.set()
            print(f"POSTED_STOP user={label} page={page} {err}", file=sys.stderr, flush=True)
            return rows, err
        upstream = posted_upstream_error(data)
        inner = data.get("data") if isinstance(data.get("data"), dict) else {}
        payload = inner.get("data") if isinstance(inner.get("data"), dict) else inner
        notes = payload.get("notes") if isinstance(payload, dict) else []
        if not isinstance(notes, list):
            notes = []
        if upstream and not notes:
            print(f"POSTED_STOP user={label} page={page} {upstream}", file=sys.stderr, flush=True)
            return rows, upstream
        added = 0
        for item in notes:
            rec = extract_posted_note(item if isinstance(item, dict) else {}, label)
            if not rec:
                continue
            rows.append(rec)
            added += 1
            if not user_id and rec.get("user_id"):
                user_id = str(rec["user_id"])
        last = notes[-1] if notes and isinstance(notes[-1], dict) else {}
        cursor = str(last.get("cursor") or last.get("id") or last.get("note_id") or "").strip()
        has_more = payload.get("has_more") if isinstance(payload, dict) else None
        print(
            f"page user={label} p={page} got={len(notes)} keep={added} has_more={has_more}",
            file=sys.stderr,
            flush=True,
        )
        if not notes or has_more is False or not cursor:
            break
    return rows, None


def default_posted_filename(targets: list[dict[str, str]]) -> str:
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    first = (targets[0]["label"] if targets else "user")[:20]
    slug = re.sub(r"[\\/:*?\"<>|]", "_", first)
    extra = f"等{len(targets)}账号" if len(targets) > 1 else ""
    return f"小红书博主笔记-无正文无粉丝-{slug}{extra}-{stamp}.xlsx"


def run_user_posted_export(params: dict[str, Any]) -> dict[str, Any]:
    targets = parse_user_targets(params)
    if not targets:
        raise XhsError("必须提供 user_id / user_ids（博主账号 id，或主页/分享链接）")
    try:
        pages = int(params.get("pages") or DEFAULT_PAGES)
    except (TypeError, ValueError) as err:
        raise XhsError("pages 必须是整数") from err
    if pages < 1 or pages > MAX_POSTED_PAGES:
        raise XhsError(f"pages 必须在 1–{MAX_POSTED_PAGES} 之间")
    try:
        workers = int(params.get("workers") or DEFAULT_WORKERS)
    except (TypeError, ValueError) as err:
        raise XhsError("workers 必须是整数") from err
    workers = max(1, min(workers, MAX_WORKERS, DEFAULT_QPS))
    export_dir = Path(params.get("export_dir") or "./workspace")
    filename = ensure_xlsx_filename(
        str(params.get("filename") or "").strip() or default_posted_filename(targets)
    )
    if "/" in filename or "\\" in filename:
        raise XhsError("filename 只能是文件名，不能含路径")

    proxy = optional_env("SKILL_TIKHUB_PROXY")
    api = TikHub(
        require_env("SKILL_TIKHUB_API_KEY"),
        require_env("SKILL_TIKHUB_BASE_URL"),
        DEFAULT_QPS,
        proxy=proxy,
    )
    try:
        print(
            f"users={len(targets)} pages={pages} workers={workers} with_body=False with_fans=False"
            + (f" proxy={proxy_host(proxy)}" if proxy else " proxy=off"),
            file=sys.stderr,
            flush=True,
        )
        batches: list[list[dict[str, Any]]] = []
        per_user: dict[str, int] = {}
        warnings: list[str] = []
        stop = threading.Event()
        with ThreadPoolExecutor(max_workers=min(workers, len(targets))) as pool:
            futs = {
                pool.submit(fetch_user_posted, api, target, pages, stop): target for target in targets
            }
            for fut in as_completed(futs):
                target = futs[fut]
                label = target["label"]
                try:
                    batch, err = fut.result()
                except Exception as err:  # noqa: BLE001
                    print(f"USER_FAIL {label} {err}", file=sys.stderr, flush=True)
                    warnings.append(f"{label}: {err}")
                    batch, err = [], str(err)
                per_user[label] = len(batch)
                batches.append(batch)
                if err:
                    kept = f"，已保留前 {len(batch)} 条" if batch else ""
                    warnings.append(f"{label}: {err}{kept}")
        merged = merge_rows(batches)
        raw_total = sum(len(batch) for batch in batches)
        print(f"DEDUP {raw_total} -> {len(merged)}", file=sys.stderr, flush=True)
        if raw_total == 0:
            raise XhsError("；".join(uniq_warnings(warnings)) or "未获取到笔记，请确认博主账号 id 有效")
        out_path = (export_dir / filename).resolve()
        write_excel(out_path, merged, {}, with_body=False, with_fans=False)
        cost = build_cost_from_paths(
            api.billed,
            [(USER_POSTED_PATH, "get_user_posted_notes 翻页")],
        )
    finally:
        api.close()

    types = Counter(media_type(str(rec["type"])) for rec in merged.values())
    warn_list = uniq_warnings(warnings)
    note = "请点击下载按钮获取 Excel。"
    if warn_list:
        note = (
            f"中途中断（常见原因：TikHub 付费余额不足），已导出已获取的 {len(merged)} 条。"
            f"{'；'.join(warn_list)} 请点击下载按钮获取 Excel。"
        )
    note += f" 本次消耗约 ${cost['total_usd']:.2f}，回复中须贴出 data.cost.table。"
    return {
        "output_xlsx": out_path.name,
        "output_path": str(out_path),
        "user_ids": [t["label"] for t in targets],
        "pages": pages,
        "per_user": per_user,
        "raw_rows": raw_total,
        "rows": len(merged),
        "types": dict(types),
        "detail_ok": 0,
        "detail_fail": 0,
        "empty_title": sum(1 for rec in merged.values() if not rec["title"]),
        "with_fans": False,
        "with_body": False,
        "partial": bool(warn_list),
        "warnings": warn_list,
        "cost": cost,
        "note": note,
    }


def emit_user_posted_main() -> None:
    raw = sys.argv[1] if len(sys.argv) > 1 else ""
    try:
        params = json.loads(raw) if str(raw).strip() else {}
        if not isinstance(params, dict):
            raise XhsError("参数必须是 JSON 对象")
        data = run_user_posted_export(params)
        print(json.dumps({"ok": True, "data": data}, ensure_ascii=False))
    except XhsError as err:
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": str(err),
                    "hint": "提供 user_id；确认 SKILL_TIKHUB_API_KEY / SKILL_TIKHUB_BASE_URL；59 需 SKILL_TIKHUB_PROXY；错误 user_id 仍会计费；402 表示该路由要付费余额",
                },
                ensure_ascii=False,
            )
        )
    except Exception as err:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": f"{type(err).__name__}: {err}"}, ensure_ascii=False))
