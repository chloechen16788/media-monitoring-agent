#!/usr/bin/env python3
"""Format weekly-news Excel: GPT(OpenCode)+DeepSeek summaries, DeepSeek tags & patents.

Ported from the Cursor skill `news-excel-format` to the codex-agent sub skill
contract (JSON in/out, ./workspace output, env-only secrets). The OpenCode call
that previously shelled out to opencode-go-api/call.py is inlined here.
"""

from __future__ import annotations

import json
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook
from openpyxl.styles import Alignment

SUMMARY_PROMPT = (
    "你是资讯分析师，任务：1. 读懂科技新闻原文，摘要仅依托原文现有信息，禁止拓展、编造原文不存在的内容；"
    "2. 从原文提取主体/核心事件 / 产品，仅总结文本内写明的主体背景，科技价值、潜在影响。原文无相关描述则不写主体背景；"
    "3. 输出单段式摘要，全文不超 250 字；4. 若原文正文内容极少（仅单句 / 简短几句话），摘要严格贴合原文篇幅精简输出，只复述原文现有信息，绝不自行扩充篇幅、添加额外观点与延伸内容；"
    "5. 摘要不提及日期.全程中文输出，直接给出摘要结果，无需多余内容。"
)
TAG_PROMPT = (
    "请从摘要抽取分类和主体，行业包括汽车新闻（包含智驾，传感器，芯片，座舱，电池，三电，底盘，车联网，工业制造等），"
    "公司经营（公司计划，战略，合作等），AI新闻，元宇宙新闻（AR，VR ，眼镜相关），新能源新闻，新材料新闻，"
    "机器人新闻（机器人相关算法和产品），其他科技新闻；主体包括公司和机构具体名称，不超过3个，可用简称如：百度，蔚来。"
    "命中多个分类请用逗号分隔\n注意：\n1.其他科技新闻与其他分类互斥\n"
    "2.必须返回摘要中实际存在的主体名称，仅输出标签，不要其他内容。返回如：元宇宙新闻,索尼"
)
PATENT_PROMPT = "查看内容中是否有专利号，包括申请号和公开号，取出来，只显示专利号不要别的，多个号逗号分隔，没有显示无"

TITLE_COL = "新闻标题"
BODY_COL = "新闻正文"
GPT_COL = "正文摘要-GPT"
DS_COL = "正文摘要-Deepseek"
TAG_COL = "备注"
PATENT_COL = "专利号"
WRITE_COLS = [GPT_COL, DS_COL, TAG_COL, PATENT_COL]

OPENCODE_DEFAULT_BASE = "https://opencode.ai/zen/go/v1"
OPENCODE_DEFAULT_MODEL = "gpt-5.6-luna"
RESPONSES_MODELS = {"gpt-5.6-luna", "grok-4.5"}
MAX_RETRIES = 2


class FormatError(Exception):
    pass


def require_keys() -> None:
    if not (os.environ.get("OPENCODE_API_KEY") or "").strip():
        raise FormatError("缺少 OPENCODE_API_KEY 环境变量")
    if not (os.environ.get("DEEPSEEK_API_KEY") or "").strip():
        raise FormatError("缺少 DEEPSEEK_API_KEY 环境变量")


# ---- DeepSeek (JSON object) -------------------------------------------------

def ds_one(prompt: str, index: int) -> dict:
    url = os.environ.get("DEEPSEEK_BASE_URL", "https://api.deepseek.com").rstrip("/") + "/chat/completions"
    model = os.environ.get("DEEPSEEK_MODEL", "deepseek-v4-flash")
    body = {
        "model": model,
        "messages": [{"role": "user", "content": prompt}],
        "response_format": {"type": "json_object"},
        "extra_body": {"thinking": {"type": "disabled"}},
    }
    headers = {"Content-Type": "application/json", "Authorization": f"Bearer {os.environ['DEEPSEEK_API_KEY']}"}
    last_error = ""
    for attempt in range(MAX_RETRIES + 1):
        try:
            req = Request(url, method="POST", headers=headers, data=json.dumps(body, ensure_ascii=False).encode("utf-8"))
            with urlopen(req, timeout=90) as resp:
                data = json.loads(resp.read().decode("utf-8", errors="replace"))
            parsed = json.loads(data["choices"][0]["message"]["content"])
            return {"index": index, "ok": True, "data": parsed, "error": None}
        except HTTPError as err:
            last_error = f"HTTP {err.code}: {err.read().decode('utf-8', errors='replace')[:300]}"
        except Exception as err:  # noqa: BLE001
            last_error = str(err)
        if attempt < MAX_RETRIES:
            time.sleep(1.5 * (2**attempt))
    return {"index": index, "ok": False, "data": {}, "error": last_error}


def ds_batch(prompts: list[str], workers: int = 8) -> list[dict]:
    results: list[dict] = []
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        futs = {pool.submit(ds_one, prompt, i): i for i, prompt in enumerate(prompts)}
        for fut in as_completed(futs):
            results.append(fut.result())
    results.sort(key=lambda item: item["index"])
    return results


# ---- OpenCode (inlined) -----------------------------------------------------

def _oc_extract_text(api: str, body) -> str | None:
    if not isinstance(body, dict):
        return None
    if api == "chat":
        choices = body.get("choices") or []
        if choices:
            return (choices[0].get("message") or {}).get("content")
        return None
    texts: list[str] = []
    if body.get("output_text"):
        texts.append(str(body["output_text"]))
    for item in body.get("output") or []:
        if not isinstance(item, dict):
            continue
        if item.get("text"):
            texts.append(str(item["text"]))
        for part in item.get("content") or []:
            if isinstance(part, dict) and part.get("text"):
                texts.append(str(part["text"]))
    return "\n".join(texts) if texts else None


def oc_one(base: str, key: str, model: str, api: str, prompt: str, index: int) -> dict:
    if api == "responses":
        url = f"{base}/responses"
        payload = {"model": model, "input": prompt}
    else:
        url = f"{base}/chat/completions"
        payload = {"model": model, "messages": [{"role": "user", "content": prompt}], "max_tokens": 1024}
    req = Request(
        url,
        data=json.dumps(payload).encode("utf-8"),
        method="POST",
        headers={
            "Authorization": f"Bearer {key}",
            "Accept": "application/json",
            "Content-Type": "application/json",
            "User-Agent": "curl/8.7.1",
        },
    )
    try:
        with urlopen(req, timeout=90) as resp:
            body = json.loads(resp.read().decode("utf-8", errors="replace") or "{}")
        text = _oc_extract_text(api, body)
        return {"index": index, "ok": bool(text), "text": (text or "").strip()}
    except (HTTPError, URLError, TimeoutError, json.JSONDecodeError) as err:  # noqa: BLE001
        return {"index": index, "ok": False, "text": "", "error": str(err)}


def gpt_summaries(prompts: list[str]) -> list[str]:
    key = os.environ["OPENCODE_API_KEY"].strip()
    base = os.environ.get("OPENCODE_BASE_URL", OPENCODE_DEFAULT_BASE).rstrip("/")
    model = os.environ.get("OPENCODE_MODEL", OPENCODE_DEFAULT_MODEL)
    api = "responses" if model in RESPONSES_MODELS else "chat"
    concurrency = max(1, int(os.environ.get("OPENCODE_CONCURRENCY", "30")))
    texts = [""] * len(prompts)
    with ThreadPoolExecutor(max_workers=min(concurrency, max(1, len(prompts)))) as pool:
        futs = {pool.submit(oc_one, base, key, model, api, prompt, i): i for i, prompt in enumerate(prompts)}
        for fut in as_completed(futs):
            res = fut.result()
            texts[res["index"]] = (res.get("text") or "").strip()
    return texts


# ---- Excel ------------------------------------------------------------------

def header_map(ws) -> dict[str, int]:
    mapping = {}
    for idx in range(1, ws.max_column + 1):
        value = ws.cell(1, idx).value
        if value is not None and str(value).strip():
            mapping[str(value).strip()] = idx
    return mapping


def ensure_columns(ws, mapping: dict[str, int]) -> dict[str, int]:
    for name in WRITE_COLS:
        if name in mapping:
            continue
        col = ws.max_column + 1
        ws.cell(1, col, name)
        mapping[name] = col
    return mapping


def read_rows(ws) -> list[dict]:
    mapping = header_map(ws)
    if TITLE_COL not in mapping or BODY_COL not in mapping:
        raise FormatError(f"缺少列 {TITLE_COL}/{BODY_COL}，当前表头={list(mapping)}")
    rows = []
    for excel_row in range(2, ws.max_row + 1):
        title = ws.cell(excel_row, mapping[TITLE_COL]).value
        body = ws.cell(excel_row, mapping[BODY_COL]).value
        if title is None and body is None:
            continue
        rows.append({
            "excel_row": excel_row,
            "title": "" if title is None else str(title).strip(),
            "body": "" if body is None else str(body).strip(),
        })
    return rows


def run(params: dict) -> dict:
    input_path = params.get("input_path") or params.get("input")
    if not input_path:
        raise FormatError("必须提供 input_path（待格式化 Excel）")
    src = Path(str(input_path)).expanduser()
    if not src.is_file():
        raise FormatError(f"输入文件不存在: {src}")
    sheet = params.get("sheet")
    export_dir = Path(params.get("export_dir") or "./workspace")

    require_keys()
    wb = load_workbook(src)
    sheet_name = sheet or wb.sheetnames[0]
    if sheet_name not in wb.sheetnames:
        raise FormatError(f"sheet 不存在: {sheet_name}，可选={wb.sheetnames}")
    ws = wb[sheet_name]
    rows = read_rows(ws)
    if not rows:
        raise FormatError(f"{sheet_name} 无数据行")

    summary_prompts = [
        f"{SUMMARY_PROMPT}\n\n【新闻标题】\n{row['title']}\n\n【新闻正文】\n{row['body']}" for row in rows
    ]
    patent_prompts = [
        f"{PATENT_PROMPT}\n\n请只根据下面正文抽取，输出 JSON：{{\"patent\":\"专利号或无\"}}\n\n【新闻标题】\n{row['title']}\n\n【新闻正文】\n{row['body']}"
        for row in rows
    ]
    ds_summary_prompts = [
        prompt + "\n\n请把摘要放进 JSON：{\"summary\":\"单段中文摘要\"}，不要其他字段。"
        for prompt in summary_prompts
    ]

    n = len(rows)
    print(f"开始格式化：共 {n} 行", file=sys.stderr, flush=True)
    print(f"[1/4] GPT 正文摘要 {n} 行…", file=sys.stderr, flush=True)
    gpt_texts = gpt_summaries(summary_prompts)
    print(f"[2/4] DeepSeek 正文摘要 {n} 行…", file=sys.stderr, flush=True)
    ds_sum = ds_batch(ds_summary_prompts)
    print(f"[3/4] DeepSeek 专利号抽取 {n} 行…", file=sys.stderr, flush=True)
    ds_pat = ds_batch(patent_prompts)

    tag_prompts = []
    ds_texts = []
    for i, row in enumerate(rows):
        summary = str((ds_sum[i].get("data") or {}).get("summary") or "").strip()
        ds_texts.append(summary)
        tag_prompts.append(
            f"{TAG_PROMPT}\n\n请输出 JSON：{{\"label\":\"分类,主体\"}}\n\n【新闻标题】\n{row['title']}\n\n【摘要】\n{summary}"
        )
    print(f"[4/4] DeepSeek 分类打标 {n} 行…", file=sys.stderr, flush=True)
    ds_tags = ds_batch(tag_prompts)

    mapping = ensure_columns(ws, header_map(ws))
    wrap = Alignment(wrap_text=True, vertical="top")
    patent_hit = 0
    for i, row in enumerate(rows):
        patent = str((ds_pat[i].get("data") or {}).get("patent") or "无").strip() or "无"
        label = str((ds_tags[i].get("data") or {}).get("label") or "").strip()
        if patent not in ("无", "没有"):
            patent_hit += 1
        values = {GPT_COL: gpt_texts[i], DS_COL: ds_texts[i], TAG_COL: label, PATENT_COL: patent}
        for name, value in values.items():
            cell = ws.cell(row["excel_row"], mapping[name])
            cell.value = value
            cell.alignment = wrap

    export_dir.mkdir(parents=True, exist_ok=True)
    out_path = (export_dir / f"{src.stem}_格式化.xlsx").resolve()
    wb.save(out_path)

    return {
        "output_path": str(out_path),
        "sheet": sheet_name,
        "rows": len(rows),
        "patent_hit": patent_hit,
        "gpt_ok": sum(1 for t in gpt_texts if t),
        "ds_ok": sum(1 for t in ds_texts if t),
    }


def main() -> None:
    try:
        args = json.loads(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1] else {}
        data = run(args if isinstance(args, dict) else {})
        print(json.dumps({"ok": True, "data": data}, ensure_ascii=False))
    except FormatError as err:
        print(json.dumps({"ok": False, "error": str(err), "hint": "确认 input_path 与 OPENCODE_API_KEY/DEEPSEEK_API_KEY 环境变量"}, ensure_ascii=False))
    except Exception as err:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": f"{type(err).__name__}: {err}"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
