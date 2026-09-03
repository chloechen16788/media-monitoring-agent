#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Reddit keyword search via TikHub APP APIs; export posts + comments Excel."""

from __future__ import annotations

import json
import os
import re
import sys
import threading
import time
from collections import Counter, OrderedDict, deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
)
SEARCH_PATH = "/api/v1/reddit/app/fetch_dynamic_search"
DETAIL_PATH = "/api/v1/reddit/app/fetch_post_details_batch_large"
COMMENTS_PATH = "/api/v1/reddit/app/fetch_post_comments"
UNIT_USD = 0.001
DEFAULT_QPS = 10
DEFAULT_WORKERS = 5
DEFAULT_PAGES = 1
MAX_PAGES = 20
DEFAULT_COMMENT_PAGES = 2
MAX_COMMENT_PAGES = 5
MAX_WORKERS = 10
BATCH_SIZE = 30
SPLIT_RE = re.compile(r"[\n,，;；]+")
HEADERS = ["日期", "标题", "内容", "链接", "作者", "类型", "互动数"]
SORT_ALIASES = {
    "new": "NEW",
    "最新": "NEW",
    "时间": "NEW",
    "时间排序": "NEW",
    "relevance": "RELEVANCE",
    "相关": "RELEVANCE",
    "相关性": "RELEVANCE",
    "hot": "HOT",
    "热门": "HOT",
    "top": "TOP",
    "最受欢迎": "TOP",
    "comments": "COMMENTS",
    "评论数": "COMMENTS",
    "评论": "COMMENTS",
}
SORT_OK = {"RELEVANCE", "HOT", "TOP", "NEW", "COMMENTS"}
TIME_ALIASES = {
    "all": "all",
    "全部": "all",
    "所有时间": "all",
    "year": "year",
    "年": "year",
    "去年": "year",
    "month": "month",
    "月": "month",
    "上个月": "month",
    "近一个月": "month",
    "week": "week",
    "周": "week",
    "上周": "week",
    "day": "day",
    "今天": "day",
    "日": "day",
    "hour": "hour",
    "小时": "hour",
    "过去1小时": "hour",
}
TIME_OK = {"all", "year", "month", "week", "day", "hour"}


class RedditError(Exception):
    pass


class RateLimiter:
    def __init__(self, qps: int) -> None:
        self.qps = max(1, qps)
        self._hits: deque[float] = deque()
        self._lock = threading.Lock()

    def acquire(self) -> None:
        while True:
            with self._lock:
                now = time.monotonic()
                while self._hits and now - self._hits[0] >= 1.0:
                    self._hits.popleft()
                if len(self._hits) < self.qps:
                    self._hits.append(now)
                    return
                wait = 1.0 - (now - self._hits[0])
            time.sleep(max(wait, 0.02))


def require_env(name: str) -> str:
    value = (os.environ.get(name) or "").strip()
    if not value:
        raise RedditError(f"missing env: {name}")
    return value


def optional_env(name: str) -> str:
    return (os.environ.get(name) or "").strip()


def proxy_host(proxy_url: str) -> str:
    rest = proxy_url.split("://", 1)[-1]
    return rest.split("@")[-1].split("/")[0] or "on"


def tikhub_error_text(resp: httpx.Response) -> str:
    try:
        data = resp.json()
    except Exception:  # noqa: BLE001
        return f"HTTP {resp.status_code}"
    detail = data.get("detail") if isinstance(data, dict) and isinstance(data.get("detail"), dict) else {}
    msg = (
        (detail.get("message_zh") if isinstance(detail, dict) else None)
        or (detail.get("message") if isinstance(detail, dict) else None)
        or (data.get("message_zh") if isinstance(data, dict) else None)
        or (data.get("message") if isinstance(data, dict) else None)
        or (data.get("msg") if isinstance(data, dict) else None)
        or resp.text[:200]
    )
    code = (detail.get("code") if isinstance(detail, dict) else None) or (
        data.get("code") if isinstance(data, dict) else None
    ) or resp.status_code
    return f"TikHub {code}: {msg}"


def is_fatal_tikhub(err: Exception | str) -> bool:
    text = str(err)
    return any(code in text for code in ("401", "402", "403"))


def parse_keywords(raw: Any) -> list[str]:
    parts: list[str] = []
    if isinstance(raw, list):
        for item in raw:
            parts.extend(SPLIT_RE.split(str(item or "")))
    elif isinstance(raw, str):
        parts.extend(SPLIT_RE.split(raw))
    out: list[str] = []
    seen: set[str] = set()
    for part in parts:
        kw = part.strip()
        if kw and kw not in seen:
            seen.add(kw)
            out.append(kw)
    return out


def normalize_sort(raw: Any) -> str:
    text = str(raw or "NEW").strip()
    mapped = SORT_ALIASES.get(text.lower(), text.upper())
    if mapped not in SORT_OK:
        raise RedditError(f"sort 仅支持 {', '.join(sorted(SORT_OK))}")
    return mapped


def normalize_time_range(raw: Any) -> str:
    text = str(raw or "month").strip()
    mapped = TIME_ALIASES.get(text.lower(), text.lower())
    if mapped not in TIME_OK:
        raise RedditError(f"time_range 仅支持 {', '.join(sorted(TIME_OK))}")
    return mapped


def ensure_xlsx_filename(name: str) -> str:
    stem = Path(str(name or "").strip()).name
    if not stem:
        return stem
    lower = stem.lower()
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        return stem
    if lower.endswith(".csv") or lower.endswith(".xls"):
        return f"{stem.rsplit('.', 1)[0]}.xlsx"
    return f"{stem}.xlsx"


def fmt_time(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        ts = int(value)
        if ts > 10**12:
            ts //= 1000
        try:
            return datetime.fromtimestamp(ts, tz=timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            return ""
    text = str(value).strip()
    if not text:
        return ""
    iso = text.replace("Z", "+00:00")
    if re.search(r"[+-]\d{4}$", iso):
        iso = iso[:-2] + ":" + iso[-2:]
    try:
        dt = datetime.fromisoformat(iso)
        if dt.tzinfo is not None:
            dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return text[:19]


def num(v: Any) -> int:
    if v is None or v == "":
        return 0
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return int(v)
    try:
        return int(float(str(v).strip().replace(",", "")))
    except Exception:
        return 0


def markdown_of(content: Any) -> str:
    if isinstance(content, dict):
        return str(content.get("markdown") or content.get("text") or "").strip()
    if isinstance(content, str):
        return content.strip()
    return ""


def author_name(obj: dict[str, Any]) -> str:
    info = obj.get("authorInfo") if isinstance(obj.get("authorInfo"), dict) else {}
    return str(info.get("name") or obj.get("author") or "").strip()


def reddit_url(raw: Any, permalink: Any = "") -> str:
    for candidate in (raw, permalink):
        s = str(candidate or "").strip()
        if not s:
            continue
        if s.startswith("http://") or s.startswith("https://"):
            return s
        if s.startswith("/"):
            return "https://www.reddit.com" + s
    return ""


def normalize_post_id(raw: Any) -> str:
    text = str(raw or "").strip()
    if not text:
        return ""
    if text.startswith("t3_"):
        return text
    if re.fullmatch(r"[A-Za-z0-9]+", text):
        return f"t3_{text}"
    return text


def unwrap_data(payload: dict[str, Any]) -> dict[str, Any]:
    data = payload.get("data")
    if isinstance(data, str):
        try:
            data = json.loads(data)
        except Exception:
            return {}
    return data if isinstance(data, dict) else {}


def uniq_warnings(items: list[str], limit: int = 6) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for item in items:
        text = str(item or "").strip()
        if not text or text in seen:
            continue
        seen.add(text)
        out.append(text)
        if len(out) >= limit:
            break
    return out


class TikHub:
    def __init__(self, key: str, base_url: str, qps: int, proxy: str = "") -> None:
        self.key = key
        self.base = base_url.rstrip("/")
        self.limiter = RateLimiter(qps)
        self.billed: Counter[str] = Counter()
        kwargs: dict[str, Any] = {
            "timeout": 60.0,
            "headers": {"Authorization": f"Bearer {key}", "User-Agent": UA},
        }
        if proxy:
            kwargs["proxy"] = proxy
        self.client = httpx.Client(**kwargs)

    def close(self) -> None:
        self.client.close()

    def get(self, path: str, params: dict[str, Any]) -> dict[str, Any]:
        clean = {k: v for k, v in params.items() if v not in (None, "")}
        last: Exception | None = None
        for attempt in range(4):
            self.limiter.acquire()
            try:
                resp = self.client.get(f"{self.base}{path}", params=clean)
                if resp.status_code in (401, 402, 403):
                    raise RedditError(tikhub_error_text(resp))
                if resp.status_code in (429, 500, 502, 503, 504):
                    last = RedditError(tikhub_error_text(resp))
                    time.sleep(1.2 * (attempt + 1))
                    continue
                if resp.status_code >= 400:
                    raise RedditError(tikhub_error_text(resp))
                data = resp.json()
                if not isinstance(data, dict):
                    raise RedditError("TikHub 返回非对象")
                self.billed[path] += 1
                return data
            except RedditError:
                raise
            except Exception as err:  # noqa: BLE001
                last = err
                time.sleep(0.8 * (attempt + 1))
        raise RedditError(f"TikHub 请求失败: {last}")


def parse_post(obj: dict[str, Any]) -> dict[str, Any] | None:
    nid = normalize_post_id(obj.get("id"))
    if not nid:
        return None
    permalink = obj.get("permalink") or ""
    return {
        "id": nid,
        "title": str(obj.get("postTitle") or obj.get("title") or "").strip(),
        "body": markdown_of(obj.get("content")),
        "url": reddit_url(obj.get("url"), permalink),
        "author": author_name(obj),
        "score": num(obj.get("score")),
        "time": fmt_time(obj.get("createdAt")),
    }


def extract_search_posts(payload: dict[str, Any]) -> tuple[list[dict[str, Any]], str, bool]:
    data = unwrap_data(payload)
    search = data.get("search") if isinstance(data.get("search"), dict) else data
    dyn = search.get("dynamic") if isinstance(search.get("dynamic"), dict) else {}
    comps = dyn.get("components") if isinstance(dyn.get("components"), dict) else {}
    main = comps.get("main") if isinstance(comps.get("main"), dict) else {}
    page_info = main.get("pageInfo") if isinstance(main.get("pageInfo"), dict) else {}
    after = str(page_info.get("endCursor") or "").strip()
    has_more = page_info.get("hasNextPage") is True
    posts: list[dict[str, Any]] = []
    for edge in main.get("edges") or []:
        node = edge.get("node") if isinstance(edge, dict) else None
        children = node.get("children") if isinstance(node, dict) else None
        if not isinstance(children, list):
            continue
        for child in children:
            post = child.get("post") if isinstance(child, dict) else None
            if isinstance(post, dict):
                rec = parse_post(post)
                if rec:
                    posts.append(rec)
    return posts, after, has_more


def search_keyword(
    api: TikHub, keyword: str, pages: int, sort: str, time_range: str, stop: threading.Event | None
) -> tuple[list[dict[str, Any]], str | None]:
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    after = ""
    for page in range(1, pages + 1):
        if stop is not None and stop.is_set():
            return rows, "TikHub 中断（常见原因：付费余额不足）"
        params: dict[str, Any] = {
            "query": keyword,
            "search_type": "post",
            "sort": sort,
            "time_range": time_range,
            "safe_search": "unset",
            "allow_nsfw": "0",
            "need_format": False,
        }
        if after:
            params["after"] = after
        try:
            data = api.get(SEARCH_PATH, params)
        except RedditError as err:
            if is_fatal_tikhub(err) and stop is not None:
                stop.set()
            print(f"SEARCH_STOP kw={keyword} page={page} {err}", file=sys.stderr, flush=True)
            return rows, str(err)
        if data.get("code") != 200:
            msg = data.get("message_zh") or data.get("message") or ""
            err = f"TikHub {data.get('code')}: {msg}".strip()
            if is_fatal_tikhub(err) and stop is not None:
                stop.set()
            return rows, err
        posts, after, has_more = extract_search_posts(data)
        added = 0
        for rec in posts:
            if rec["id"] in seen:
                continue
            seen.add(rec["id"])
            rec["keyword"] = keyword
            rows.append(rec)
            added += 1
        print(
            f"page kw={keyword} p={page} got={len(posts)} keep={added} has_more={has_more}",
            file=sys.stderr,
            flush=True,
        )
        if not posts or not has_more or not after:
            break
    return rows, None


def extract_detail_posts(payload: dict[str, Any]) -> list[dict[str, Any]]:
    data = unwrap_data(payload)
    items = data.get("postsInfoByIds")
    if not isinstance(items, list):
        return []
    out: list[dict[str, Any]] = []
    for item in items:
        if isinstance(item, dict):
            rec = parse_post(item)
            if rec:
                out.append(rec)
    return out


def fetch_details(
    api: TikHub, post_ids: list[str], stop: threading.Event | None
) -> tuple[dict[str, dict[str, Any]], list[str]]:
    found: dict[str, dict[str, Any]] = {}
    warnings: list[str] = []
    chunks = [post_ids[i : i + BATCH_SIZE] for i in range(0, len(post_ids), BATCH_SIZE)]
    for idx, chunk in enumerate(chunks, start=1):
        if stop is not None and stop.is_set():
            warnings.append("TikHub 中断（常见原因：付费余额不足）")
            break
        try:
            data = api.get(
                DETAIL_PATH,
                {
                    "post_ids": ",".join(chunk),
                    "include_comment_id": False,
                    "need_format": False,
                },
            )
        except RedditError as err:
            if is_fatal_tikhub(err) and stop is not None:
                stop.set()
            warnings.append(str(err))
            print(f"DETAIL_STOP batch={idx} {err}", file=sys.stderr, flush=True)
            break
        if data.get("code") != 200:
            msg = data.get("message_zh") or data.get("message") or ""
            err = f"TikHub {data.get('code')}: {msg}".strip()
            if is_fatal_tikhub(err) and stop is not None:
                stop.set()
            warnings.append(err)
            break
        recs = extract_detail_posts(data)
        for rec in recs:
            found[rec["id"]] = rec
        print(f"detail batch={idx}/{len(chunks)} got={len(recs)}", file=sys.stderr, flush=True)
    return found, warnings


def extract_comments(payload: dict[str, Any]) -> tuple[list[dict[str, Any]], str, bool]:
    data = unwrap_data(payload)
    info = data.get("postInfoById") if isinstance(data.get("postInfoById"), dict) else {}
    forest = info.get("commentForest") if isinstance(info.get("commentForest"), dict) else {}
    page_info = forest.get("pageInfo") if isinstance(forest.get("pageInfo"), dict) else {}
    has_more = page_info.get("hasNextPage") is True
    trees = forest.get("trees") if isinstance(forest.get("trees"), list) else []
    comments: list[dict[str, Any]] = []
    after = ""
    for tree in trees:
        if not isinstance(tree, dict):
            continue
        more = tree.get("more") if isinstance(tree.get("more"), dict) else None
        if more and more.get("cursor"):
            after = str(more.get("cursor") or "")
        node = tree.get("node") if isinstance(tree.get("node"), dict) else {}
        cid = str(node.get("id") or "").strip()
        if not cid:
            continue
        comments.append(
            {
                "id": cid,
                "body": markdown_of(node.get("content")),
                "url": reddit_url(node.get("permalink") or node.get("url")),
                "author": author_name(node),
                "score": num(node.get("score")),
                "time": fmt_time(node.get("createdAt")),
            }
        )
    return comments, after, has_more


def fetch_comments(
    api: TikHub, post_id: str, pages: int, stop: threading.Event | None
) -> tuple[list[dict[str, Any]], str | None]:
    rows: list[dict[str, Any]] = []
    seen: set[str] = set()
    after = ""
    for page in range(1, pages + 1):
        if stop is not None and stop.is_set():
            return rows, "TikHub 中断（常见原因：付费余额不足）"
        params: dict[str, Any] = {
            "post_id": post_id,
            "sort_type": "NEW",
            "need_format": False,
        }
        if after:
            params["after"] = after
        try:
            data = api.get(COMMENTS_PATH, params)
        except RedditError as err:
            if is_fatal_tikhub(err) and stop is not None:
                stop.set()
            return rows, str(err)
        if data.get("code") != 200:
            msg = data.get("message_zh") or data.get("message") or ""
            err = f"TikHub {data.get('code')}: {msg}".strip()
            if is_fatal_tikhub(err) and stop is not None:
                stop.set()
            return rows, err
        comments, after, has_more = extract_comments(data)
        for rec in comments:
            if rec["id"] in seen:
                continue
            seen.add(rec["id"])
            rows.append(rec)
        if not comments or not has_more or not after:
            break
    return rows, None


def default_filename(keywords: list[str]) -> str:
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    slug = re.sub(r"[\\/:*?\"<>|]", "_", (keywords[0] if keywords else "search")[:20])
    extra = f"等{len(keywords)}词" if len(keywords) > 1 else ""
    return f"Reddit搜索-{slug}{extra}-{stamp}.xlsx"


def write_excel(path: Path, rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reddit"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        cell.border = thin
    wrap_cols = {3, 4}
    for row in rows:
        ws.append([row.get(h, "") for h in HEADERS])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(HEADERS)):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column in wrap_cols)
            cell.border = thin
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 42
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 10
    ws.auto_filter.ref = f"A1:G{max(ws.max_row, 1)}"
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_cost(billed: Counter[str]) -> dict[str, Any]:
    labels = [
        (SEARCH_PATH, "fetch_dynamic_search 搜帖"),
        (DETAIL_PATH, "fetch_post_details_batch_large 批量详情"),
        (COMMENTS_PATH, "fetch_post_comments 评论翻页"),
    ]
    rows: list[dict[str, Any]] = []
    for path, name in labels:
        n = int(billed.get(path) or 0)
        if n <= 0:
            continue
        rows.append(
            {
                "name": name,
                "count": n,
                "unit_usd": UNIT_USD,
                "subtotal_usd": round(n * UNIT_USD, 4),
            }
        )
    total = round(sum(float(row["subtotal_usd"]) for row in rows), 4)
    lines = ["| 调用 | 次数 | 单价 | 小计 |", "|---|---:|---:|---:|"]
    for row in rows:
        lines.append(
            f"| {row['name']} | {row['count']} | ${row['unit_usd']:.3f} | **${row['subtotal_usd']:.3f}** |"
        )
    if not rows:
        lines.append("| （无已计费请求） | 0 | $0.001 | **$0.000** |")
    lines.append(f"| **合计** | | | **${total:.3f}** |")
    return {
        "unit_usd": UNIT_USD,
        "currency": "USD",
        "calls": rows,
        "total_usd": total,
        "table": "\n".join(lines),
        "note": "按 HTTP 200 计费，402/失败不扣费。Reddit APP 单价按 TikHub 公开价 $0.001/次估算。",
    }


def run_export(params: dict[str, Any]) -> dict[str, Any]:
    keywords = parse_keywords(params.get("keywords") or params.get("keyword") or params.get("query"))
    if not keywords:
        raise RedditError("必须提供 keywords / query（搜索关键词）")
    sort = normalize_sort(params.get("sort") if params.get("sort") not in (None, "") else "NEW")
    time_range = normalize_time_range(
        params.get("time_range") if params.get("time_range") not in (None, "") else "month"
    )
    try:
        pages = int(params.get("pages") if params.get("pages") not in (None, "") else DEFAULT_PAGES)
    except (TypeError, ValueError) as err:
        raise RedditError("pages 必须是整数") from err
    if pages < 1 or pages > MAX_PAGES:
        raise RedditError(f"pages 必须在 1–{MAX_PAGES} 之间")
    try:
        comment_pages = int(
            params.get("comment_pages")
            if params.get("comment_pages") not in (None, "")
            else DEFAULT_COMMENT_PAGES
        )
    except (TypeError, ValueError) as err:
        raise RedditError("comment_pages 必须是整数") from err
    if comment_pages < 1 or comment_pages > MAX_COMMENT_PAGES:
        raise RedditError(f"comment_pages 必须在 1–{MAX_COMMENT_PAGES} 之间")
    try:
        workers = int(params.get("workers") if params.get("workers") not in (None, "") else DEFAULT_WORKERS)
    except (TypeError, ValueError) as err:
        raise RedditError("workers 必须是整数") from err
    workers = max(1, min(workers, MAX_WORKERS, DEFAULT_QPS))
    export_dir = Path(params.get("export_dir") or "./workspace")
    filename = ensure_xlsx_filename(str(params.get("filename") or "").strip() or default_filename(keywords))
    if "/" in filename or "\\" in filename:
        raise RedditError("filename 只能是文件名，不能含路径")

    proxy = optional_env("SKILL_TIKHUB_PROXY")
    api = TikHub(
        require_env("SKILL_TIKHUB_API_KEY"),
        require_env("SKILL_TIKHUB_BASE_URL"),
        DEFAULT_QPS,
        proxy=proxy,
    )
    warnings: list[str] = []
    stop = threading.Event()
    posts_by_id: OrderedDict[str, dict[str, Any]] = OrderedDict()
    try:
        print(
            f"kw={len(keywords)} pages={pages} sort={sort} time_range={time_range} "
            f"comment_pages={comment_pages} workers={workers}"
            + (f" proxy={proxy_host(proxy)}" if proxy else " proxy=off"),
            file=sys.stderr,
            flush=True,
        )
        with ThreadPoolExecutor(max_workers=min(workers, len(keywords))) as pool:
            futs = {
                pool.submit(search_keyword, api, kw, pages, sort, time_range, stop): kw for kw in keywords
            }
            for fut in as_completed(futs):
                kw = futs[fut]
                try:
                    batch, err = fut.result()
                except Exception as err:  # noqa: BLE001
                    warnings.append(f"{kw}: {err}")
                    batch, err = [], str(err)
                for rec in batch:
                    posts_by_id.setdefault(rec["id"], rec)
                if err:
                    kept = f"，已保留前 {len(batch)} 条" if batch else ""
                    warnings.append(f"{kw}: {err}{kept}")
        if not posts_by_id:
            raise RedditError("；".join(uniq_warnings(warnings)) or "未搜到帖子，请换关键词或放宽 time_range")

        details, detail_warns = fetch_details(api, list(posts_by_id.keys()), stop)
        warnings.extend(detail_warns)
        for pid, rec in posts_by_id.items():
            extra = details.get(pid) or {}
            for key in ("title", "body", "url", "author", "score", "time"):
                if extra.get(key) not in (None, ""):
                    rec[key] = extra[key]

        comments_by_post: dict[str, list[dict[str, Any]]] = {pid: [] for pid in posts_by_id}
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futs = {
                pool.submit(fetch_comments, api, pid, comment_pages, stop): pid for pid in posts_by_id
            }
            done = 0
            for fut in as_completed(futs):
                pid = futs[fut]
                done += 1
                try:
                    comments, err = fut.result()
                except Exception as err:  # noqa: BLE001
                    comments, err = [], str(err)
                comments_by_post[pid] = comments
                if err:
                    warnings.append(f"{pid}: {err}")
                if done % 10 == 0 or done == len(posts_by_id):
                    print(f"comments {done}/{len(posts_by_id)}", file=sys.stderr, flush=True)

        excel_rows: list[dict[str, Any]] = []
        comment_total = 0
        for rec in posts_by_id.values():
            excel_rows.append(
                {
                    "日期": rec.get("time") or "",
                    "标题": rec.get("title") or "",
                    "内容": rec.get("body") or "",
                    "链接": rec.get("url") or "",
                    "作者": rec.get("author") or "",
                    "类型": "帖子",
                    "互动数": rec.get("score") or 0,
                }
            )
            for cmt in comments_by_post.get(rec["id"]) or []:
                comment_total += 1
                excel_rows.append(
                    {
                        "日期": cmt.get("time") or "",
                        "标题": rec.get("title") or "",
                        "内容": cmt.get("body") or "",
                        "链接": cmt.get("url") or "",
                        "作者": cmt.get("author") or "",
                        "类型": "评论",
                        "互动数": cmt.get("score") or 0,
                    }
                )
        out_path = (export_dir / filename).resolve()
        write_excel(out_path, excel_rows)
        cost = build_cost(api.billed)
    finally:
        api.close()

    warn_list = uniq_warnings(warnings)
    note = "请点击下载按钮获取 Excel。"
    if warn_list:
        note = (
            f"中途中断（常见原因：TikHub 付费余额不足），已导出已获取的 {len(posts_by_id)} 帖。"
            f"{'；'.join(warn_list)} 请点击下载按钮获取 Excel。"
        )
    note += f" 本次消耗约 ${cost['total_usd']:.3f}，回复中须贴出 data.cost.table。"
    return {
        "output_xlsx": out_path.name,
        "output_path": str(out_path),
        "keywords": keywords,
        "sort": sort,
        "time_range": time_range,
        "pages": pages,
        "comment_pages": comment_pages,
        "posts": len(posts_by_id),
        "comments": comment_total,
        "rows": len(excel_rows),
        "partial": bool(warn_list),
        "warnings": warn_list,
        "cost": cost,
        "note": note,
    }


def main() -> None:
    raw = sys.argv[1] if len(sys.argv) > 1 else ""
    try:
        params = json.loads(raw) if str(raw).strip() else {}
        if not isinstance(params, dict):
            raise RedditError("参数必须是 JSON 对象")
        if params.get("__parse_error__"):
            raise RedditError("参数不是合法 JSON")
        data = run_export(params)
        print(json.dumps({"ok": True, "data": data}, ensure_ascii=False))
    except RedditError as err:
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": str(err),
                    "hint": "提供 keywords/query；确认 SKILL_TIKHUB_API_KEY / SKILL_TIKHUB_BASE_URL；59 需 SKILL_TIKHUB_PROXY；sort 默认 NEW，time_range 默认 month；402 表示该路由要付费余额",
                },
                ensure_ascii=False,
            )
        )
    except Exception as err:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": f"{type(err).__name__}: {err}"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
