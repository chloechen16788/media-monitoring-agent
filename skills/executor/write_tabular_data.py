"""将 enriched rows 回写到 xlsx/csv/jsonl。"""

import csv
import json
import os
import re
import sys
from typing import Any

from openpyxl import load_workbook


def _emit(payload: dict[str, Any]) -> None:
    print(json.dumps(payload, ensure_ascii=False))


def fail(error: str, hint: str = "") -> None:
    payload: dict[str, Any] = {"ok": False, "error": error}
    if hint:
        payload["hint"] = hint
    _emit(payload)
    sys.exit(0)


def read_params() -> dict[str, Any]:
    raw = sys.argv[1] if len(sys.argv) > 1 else ""
    if not raw.strip():
        return {}
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        fail("参数必须是单个合法的 JSON 字符串（禁止单引号）。")
    if not isinstance(parsed, dict):
        fail("参数 JSON 必须是对象。")
    return parsed


def infer_file_type(path: str, explicit: str | None) -> str:
    if explicit:
        return explicit.lower()
    ext = os.path.splitext(path)[1].lower()
    if ext in [".xlsx", ".xlsm"]:
        return "xlsx"
    if ext == ".csv":
        return "csv"
    if ext == ".jsonl":
        return "jsonl"
    fail("无法自动识别文件类型。", "请传 file_type: xlsx/csv/jsonl。")
    return ""


def default_output_path(source_path: str, output_path: str, file_type: str) -> str:
    if output_path.strip():
        return output_path
    stem, _ = os.path.splitext(source_path)
    ext = ".xlsx" if file_type == "xlsx" else ".csv" if file_type == "csv" else ".jsonl"
    return f"{stem}_enriched{ext}"


def _sanitize_headers(raw_headers: list[Any]) -> list[str]:
    headers: list[str] = []
    used: dict[str, int] = {}
    for idx, value in enumerate(raw_headers, start=1):
        base = str(value).strip() if value is not None else ""
        base = base or f"col_{idx}"
        count = used.get(base, 0) + 1
        used[base] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return headers


def detect_update_fields(rows: list[dict[str, Any]], explicit: Any) -> list[str]:
    if explicit is not None:
        if not isinstance(explicit, list) or any(not isinstance(x, str) for x in explicit):
            fail("update_fields 必须是字符串数组。")
        return explicit

    fields: list[str] = []
    seen = set()
    for row in rows:
        for key in row.keys():
            if key.startswith("__"):
                continue
            if key not in seen:
                seen.add(key)
                fields.append(key)
    return fields


def write_xlsx(params: dict[str, Any], source_path: str, out_path: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    wb = load_workbook(source_path)
    sheet_name = params.get("sheet_name")
    sheet_index = int(params.get("sheet_index", 0))
    header_row = int(params.get("header_row", 1))
    start_row = int(params.get("start_row", header_row + 1))

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            fail(f"找不到工作表: {sheet_name}")
        ws = wb[sheet_name]
    else:
        names = wb.sheetnames
        if not names:
            fail("xlsx 文件不包含任何工作表。")
        if sheet_index < 0 or sheet_index >= len(names):
            fail(f"sheet_index 越界: {sheet_index}")
        ws = wb[names[sheet_index]]

    header_values = [ws.cell(row=header_row, column=i).value for i in range(1, ws.max_column + 1)]
    headers = _sanitize_headers(header_values)
    col_index = {name: idx + 1 for idx, name in enumerate(headers)}

    update_fields = detect_update_fields(rows, params.get("update_fields"))
    for field in update_fields:
        if field not in col_index:
            ws.cell(row=header_row, column=ws.max_column + 1, value=field)
            col_index[field] = ws.max_column

    updated_cells = 0
    for row in rows:
        source_row_no = row.get("__source_row_number")
        if isinstance(source_row_no, int) and source_row_no > 0:
            target_row = source_row_no
        else:
            row_id = row.get("__row_id")
            if not isinstance(row_id, int) or row_id <= 0:
                continue
            target_row = start_row + row_id - 1

        for field in update_fields:
            if field.startswith("__"):
                continue
            ws.cell(row=target_row, column=col_index[field], value=row.get(field, ""))
            updated_cells += 1

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    wb.save(out_path)
    return {
        "sheet_name": ws.title,
        "updated_rows": len(rows),
        "updated_fields": update_fields,
        "updated_cells": updated_cells,
    }


def write_csv_file(params: dict[str, Any], source_path: str, out_path: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    encoding = str(params.get("encoding", "utf-8-sig"))
    delimiter = str(params.get("delimiter", ","))
    update_fields = detect_update_fields(rows, params.get("update_fields"))

    with open(source_path, "r", encoding=encoding, newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        original_rows = list(reader)
        fieldnames = list(reader.fieldnames or [])

    for field in update_fields:
        if field not in fieldnames and not field.startswith("__"):
            fieldnames.append(field)

    row_index = {idx + 1: r for idx, r in enumerate(original_rows)}
    updated_cells = 0
    for row in rows:
        source_row_no = row.get("__source_row_number")
        if isinstance(source_row_no, int) and source_row_no > 1:
            data_row_no = source_row_no - 1
        else:
            row_id = row.get("__row_id")
            if not isinstance(row_id, int) or row_id <= 0:
                continue
            data_row_no = row_id

        target = row_index.get(data_row_no)
        if target is None:
            continue
        for field in update_fields:
            if field.startswith("__"):
                continue
            target[field] = row.get(field, "")
            updated_cells += 1

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    with open(out_path, "w", encoding=encoding, newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, delimiter=delimiter)
        writer.writeheader()
        for row in original_rows:
            writer.writerow(row)

    return {
        "updated_rows": len(rows),
        "updated_fields": update_fields,
        "updated_cells": updated_cells,
        "encoding": encoding,
    }


def write_jsonl_file(params: dict[str, Any], out_path: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    keep_meta = bool(params.get("keep_meta_fields", False))
    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        for row in rows:
            obj = dict(row)
            if not keep_meta:
                obj = {k: v for k, v in obj.items() if not k.startswith("__")}
            f.write(json.dumps(obj, ensure_ascii=False) + "\n")
    return {
        "updated_rows": len(rows),
        "updated_fields": detect_update_fields(rows, params.get("update_fields")),
    }


def main() -> None:
    params = read_params()
    source_path = str(params.get("source_path", "")).strip()
    if not source_path:
        fail("缺少必填参数: source_path")
    if not os.path.exists(source_path):
        fail(f"文件不存在: {source_path}")

    rows = params.get("rows")
    if not isinstance(rows, list) or any(not isinstance(x, dict) for x in rows):
        fail("缺少必填参数: rows（对象数组）。")

    file_type = infer_file_type(source_path, params.get("file_type"))
    out_path = default_output_path(source_path, str(params.get("output_path", "")), file_type)
    out_path = os.path.abspath(out_path)

    if file_type == "xlsx":
        meta = write_xlsx(params, source_path, out_path, rows)
    elif file_type == "csv":
        meta = write_csv_file(params, source_path, out_path, rows)
    elif file_type == "jsonl":
        meta = write_jsonl_file(params, out_path, rows)
    else:
        fail(f"不支持的 file_type: {file_type}", "支持 xlsx/csv/jsonl。")
        return

    _emit(
        {
            "ok": True,
            "data": {
                "source_path": os.path.abspath(source_path),
                "output_path": out_path,
                "file_type": file_type,
                "meta": meta,
            },
        }
    )


if __name__ == "__main__":
    main()
