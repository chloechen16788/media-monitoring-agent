#!/usr/bin/env python3
"""Label Excel rows (标题 + 正文) with DeepSeek, append 5 tag columns, export xlsx.

Ported from the Cursor skill `deepseek-excel-tagging` to the codex-agent sub
skill contract (JSON in/out, ./workspace output, env-only secrets).
"""

from __future__ import annotations

import json
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

OUTPUT_COLUMNS = ["is_important_tech_news", "reason", "tech_category", "location", "subject"]
MAX_RETRIES = 2
RETRY_BACKOFF_SECONDS = 1.5


class TaggingError(Exception):
    pass


def get_runtime_config() -> dict[str, str]:
    api_key = (os.getenv("DEEPSEEK_API_KEY") or "").strip()
    if not api_key:
        raise TaggingError("缺少 DEEPSEEK_API_KEY 环境变量")
    base_url = (os.getenv("DEEPSEEK_BASE_URL") or "https://api.deepseek.com").rstrip("/")
    model = os.getenv("DEEPSEEK_MODEL") or "deepseek-v4-flash"
    return {"api_key": api_key, "base_url": base_url, "model": model}


def build_prompt(title: str, content: str) -> str:
    return f"""
你是一个专业的新闻数据标注专家。请对【标题】与【正文】进行深度语义分析，判定其是否为“高价值硬核科技新闻”。

【判定核心：必须同时满足以下 4 点才标为“是”】
1. 技术领域：属于 AI大模型/芯片、自动驾驶(智驾系统/传感器硬件)、前沿硬科技(机器人/脑机/新材料/核聚变)。
2. 地域要求：必须是【国内】主体（中港澳台企业）主导，或国内团队主导的研发。排除外企在中国研发、排除合资品牌（如别克、奥迪、大众），如果对品牌所属国家不了解，标为【不确定】。
3. 性质要求：必须是【正向】的【具体技术突破】或【基于新技术的新品首发】。
4. 严格排除项（只要命中任一则标为“否”，标题占比重于正文。）：
   - 纯应用/商业落地：如“机器人进电影院卖爆米花”、“巡检机器人中标”。
   - 营销/常规更新：新车上市（重点是价格）、改款车型、新配色、内饰升级、配置堆砌（如增加了冰箱/大灯/气囊）、销量/交付/市占率数据。
   - 非技术核心：融资信息（重点在钱）、IPO信息、中标/招标、获奖/荣誉、人事变动、战略规划、行业标准/政策、法律责任讨论（伦理）。
   - 汇总类：日报、周报、多个不相关新闻的合集。
   - 快讯类：内容少于300字。

【待处理数据】
标题：{title}
正文：{content}

【输出格式】请直接输出 JSON：
{{
    "is_important_tech_news": "是/否",
    "reason": "简要说明",
    "tech_category": "分类标签",
    "location": "国内/国外/不确定",
    "subject": "主体名称"
}}
""".strip()


def normalize_label_result(raw: dict[str, Any]) -> dict[str, str]:
    result: dict[str, str] = {}
    for key in OUTPUT_COLUMNS:
        value = raw.get(key, "")
        result[key] = "" if value is None else str(value).strip()
    if not result["is_important_tech_news"]:
        result["is_important_tech_news"] = "错误"
    return result


def call_deepseek_tag(cfg: dict[str, str], title: str, content: str, row_idx: int) -> dict[str, str]:
    url = f"{cfg['base_url']}/chat/completions"
    prompt = build_prompt(title, content)
    body = {
        "model": cfg["model"],
        "messages": [{"role": "user", "content": prompt}],
        "response_format": {"type": "json_object"},
        "extra_body": {"thinking": {"type": "disabled"}},
    }
    headers = {"Content-Type": "application/json", "Authorization": f"Bearer {cfg['api_key']}"}
    last_error = ""
    for attempt in range(MAX_RETRIES + 1):
        try:
            req = Request(url, method="POST", headers=headers, data=json.dumps(body, ensure_ascii=False).encode("utf-8"))
            with urlopen(req, timeout=60) as resp:
                data = json.loads(resp.read().decode("utf-8", errors="ignore"))
        except HTTPError as err:
            detail = err.read().decode("utf-8", errors="ignore")
            last_error = f"HTTP {err.code}: {detail[:500]}"
            if err.code in (429, 500, 502, 503, 504) and attempt < MAX_RETRIES:
                time.sleep(RETRY_BACKOFF_SECONDS * (2**attempt))
                continue
            return normalize_label_result({"is_important_tech_news": "错误", "reason": last_error})
        except Exception as err:  # noqa: BLE001
            last_error = str(err)
            if attempt < MAX_RETRIES:
                time.sleep(RETRY_BACKOFF_SECONDS * (2**attempt))
                continue
            return normalize_label_result({"is_important_tech_news": "错误", "reason": last_error})
        try:
            parsed = json.loads(data["choices"][0]["message"]["content"])
            return normalize_label_result(parsed)
        except Exception as err:  # noqa: BLE001
            last_error = f"解析失败: {err}"
            if attempt < MAX_RETRIES:
                time.sleep(RETRY_BACKOFF_SECONDS * (2**attempt))
                continue
            return normalize_label_result({"is_important_tech_news": "错误", "reason": last_error})
    return normalize_label_result({"is_important_tech_news": "错误", "reason": last_error or f"row={row_idx}"})


def run(params: dict) -> dict:
    input_path = params.get("input_path") or params.get("input")
    if not input_path:
        raise TaggingError("必须提供 input_path（待打标 Excel）")
    src = Path(str(input_path)).expanduser()
    if not src.is_file():
        raise TaggingError(f"输入文件不存在: {src}")
    sheet = params.get("sheet")
    title_col = params.get("title_col", "标题")
    content_col = params.get("content_col", "正文")
    max_workers = max(1, int(params.get("max_workers", 4)))
    export_dir = Path(params.get("export_dir") or "./workspace")

    cfg = get_runtime_config()
    wb = load_workbook(src)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]

    header_row = 1
    headers = [ws.cell(row=header_row, column=i).value for i in range(1, ws.max_column + 1)]
    header_map = {str(v).strip(): idx + 1 for idx, v in enumerate(headers) if v is not None}
    if title_col not in header_map or content_col not in header_map:
        raise TaggingError(
            f"找不到列：title={title_col}, content={content_col}，当前表头={list(header_map.keys())}"
        )
    title_col_idx = header_map[title_col]
    content_col_idx = header_map[content_col]

    out_start_col = ws.max_column + 1
    for offset, col_name in enumerate(OUTPUT_COLUMNS):
        ws.cell(row=header_row, column=out_start_col + offset, value=col_name)

    tasks: list[tuple[int, str, str]] = []
    for row_idx in range(2, ws.max_row + 1):
        title = ws.cell(row=row_idx, column=title_col_idx).value
        content = ws.cell(row=row_idx, column=content_col_idx).value
        tasks.append((
            row_idx,
            "" if title is None else str(title).strip(),
            "" if content is None else str(content).strip(),
        ))

    total = len(tasks)
    print(f"开始打标：共 {total} 行，并发 {max_workers}", file=sys.stderr, flush=True)
    results: dict[int, dict[str, str]] = {}
    done = 0
    with ThreadPoolExecutor(max_workers=max_workers) as pool:
        fut_map = {
            pool.submit(call_deepseek_tag, cfg, title, content, row_idx): row_idx
            for row_idx, title, content in tasks
        }
        for fut in as_completed(fut_map):
            row_idx = fut_map[fut]
            try:
                results[row_idx] = fut.result()
            except Exception as err:  # noqa: BLE001
                results[row_idx] = normalize_label_result({"is_important_tech_news": "错误", "reason": str(err)})
            done += 1
            if done == total or done % 5 == 0:
                print(f"打标进度 {done}/{total}", file=sys.stderr, flush=True)

    for row_idx, labels in results.items():
        for offset, col_name in enumerate(OUTPUT_COLUMNS):
            ws.cell(row=row_idx, column=out_start_col + offset, value=labels.get(col_name, ""))

    export_dir.mkdir(parents=True, exist_ok=True)
    out_path = (export_dir / f"{src.stem}_tagged.xlsx").resolve()
    wb.save(out_path)

    important = sum(1 for r in results.values() if r.get("is_important_tech_news") == "是")
    errors = sum(1 for r in results.values() if r.get("is_important_tech_news") == "错误")
    return {
        "output_path": str(out_path),
        "rows": len(tasks),
        "important": important,
        "errors": errors,
        "columns_added": OUTPUT_COLUMNS,
    }


def main() -> None:
    try:
        args = json.loads(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1] else {}
        data = run(args if isinstance(args, dict) else {})
        print(json.dumps({"ok": True, "data": data}, ensure_ascii=False))
    except TaggingError as err:
        print(json.dumps({"ok": False, "error": str(err), "hint": "确认 input_path/列名 与 DEEPSEEK_API_KEY 环境变量"}, ensure_ascii=False))
    except Exception as err:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": f"{type(err).__name__}: {err}"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
