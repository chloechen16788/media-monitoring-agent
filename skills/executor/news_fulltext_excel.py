#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Collect news full text from a title+url list and export a fixed 7-column Excel.

Ported from the Cursor skill `news-fulltext-excel` to the codex-agent sub skill
contract: pass items=[{title,url}] (and optional patches) as JSON args; output
xlsx is written under ./workspace. No API keys required.
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
from datetime import date
from html import unescape
from pathlib import Path
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)
SHORT_BODY = 80
HEADERS = ["序号", "标题", "作者", "媒体名称", "url", "日期", "正文"]
CUT_MARKERS = [
    "相关推荐", "相关新闻", "相关文章", "猜你想看", "海量资讯", "特别声明",
    "内容由作者提供", "标签:", "值班编辑", "【纠错】", "【责任编辑",
    "打开腾讯新闻", "打开凤凰新闻", "打开网易新闻", "返回搜狐", "平台声明", "网易跟贴",
]


class FulltextError(Exception):
    pass


def fetch_html(url: str, dest: Path) -> int:
    cmd = [
        "curl", "-sL", "--max-time", "25", "-A", UA,
        "-H", "Accept: text/html,application/xhtml+xml",
        "-H", "Accept-Language: zh-CN,zh;q=0.9,en;q=0.8",
        "-o", str(dest), "-w", "%{http_code}", url,
    ]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, check=False)
        return int((proc.stdout or "0").strip() or "0")
    except Exception:
        return 0


def decode_html(path: Path) -> str:
    raw = path.read_bytes()
    head = raw[:4000].decode("latin1", errors="ignore").lower()
    if any(x in head for x in ("charset=gbk", "charset=gb2312", "charset=gb18030")):
        return raw.decode("gb18030", errors="replace")
    text = raw.decode("utf-8", errors="replace")
    if text.count("\ufffd") > 50:
        return raw.decode("gb18030", errors="replace")
    return text


def meta(soup: BeautifulSoup, *names: str) -> str:
    for n in names:
        tag = soup.find("meta", attrs={"name": n}) or soup.find("meta", attrs={"property": n})
        if tag and tag.get("content"):
            return str(tag["content"]).strip()
    return ""


def first_re(text: str, pattern: str, group: int = 1) -> str:
    m = re.search(pattern, text, re.I)
    return m.group(group).strip() if m else ""


def paragraph_text(el) -> str:
    if el is None:
        return ""
    blocks = []
    for child in el.children:
        name = getattr(child, "name", None)
        if name in {"script", "style"}:
            continue
        t = str(child).strip() if name is None else child.get_text(" ", strip=True)
        if t:
            blocks.append(t)
    if blocks:
        return "\n".join(blocks)
    return el.get_text(" ", strip=True)


def clean_text(text: str) -> str:
    text = unescape(text or "")
    text = text.replace("\xa0", " ").replace("\u3000", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    lines: list[str] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            if lines and lines[-1] != "":
                lines.append("")
            continue
        lines.append(line)
    while lines and lines[0] == "":
        lines.pop(0)
    while lines and lines[-1] == "":
        lines.pop()
    return "\n".join(lines)


def cut_after(text: str, markers: list[str] | None = None) -> str:
    markers = markers or CUT_MARKERS
    for m in markers:
        idx = text.find(m)
        if idx > 80:
            text = text[:idx].rstrip()
    return text


def dedupe_body(text: str) -> str:
    t = text.strip()
    if len(t) < 400:
        return t
    paras = [p.strip() for p in re.split(r"\n{2,}", t) if p.strip()]
    out: list[str] = []
    for p in paras:
        if out and p == out[-1]:
            continue
        out.append(p)
    if len(out) >= 4 and len(out) % 2 == 0:
        mid = len(out) // 2
        if out[:mid] == out[mid:]:
            out = out[:mid]
    return "\n\n".join(out)


def normalize_date(s: str) -> str:
    s = (s or "").strip()
    s = s.replace("年", "-").replace("月", "-").replace("日", " ")
    s = s.replace("/", "-").replace("T", " ").replace("+08:00", "")
    s = re.sub(r"\s+", " ", s).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:\s+(\d{1,2}:\d{2}(?::\d{2})?))?", s)
    if not m:
        return s
    out = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    if m.group(4):
        parts = m.group(4).split(":")
        out += " " + ":".join(f"{int(p):02d}" for p in parts)
    return out


def pick_first(*vals: str) -> str:
    for v in vals:
        if v and str(v).strip():
            return str(v).strip()
    return ""


def host_of(url: str) -> str:
    return (urlparse(url).hostname or "").lower()


def generic_body(soup: BeautifulSoup) -> str:
    selectors = [
        "#artibody", "#paragraph", "#detail", "#mp-editor", "#content1",
        ".rich_media_content", ".post_body", ".v_news_content", ".contentStr",
        "section.article_content", ".article-content", ".TRS_Editor", "article",
        "#content", ".article",
    ]
    best = ""
    for sel in selectors:
        el = soup.select_one(sel)
        if not el:
            continue
        t = paragraph_text(el)
        if len(t) > len(best):
            best = t
    return best


def generic_fields(soup: BeautifulSoup, html: str) -> tuple[str, str, str]:
    author = pick_first(
        meta(soup, "author"),
        first_re(html, r"作者[：:]\s*(?:<[^>]+>)*\s*([^<\n]{1,40})"),
        first_re(html, r"本报讯（记者\s*([^）]+)）"),
        first_re(html, r"记者[：:\s]+([^\s，。\|]{2,20})"),
    )
    if author in {"网易", "腾讯网", "新浪"}:
        author = ""
    media = pick_first(
        meta(soup, "mediaid", "source", "article:author", "og:site_name"),
        first_re(html, r"来源[：:]\s*(?:<[^>]+>)*\s*([^<\n]{1,40})"),
    )
    dt = pick_first(
        meta(soup, "article:published_time", "publishdate", "og:release_date", "bytedance:published_time"),
        first_re(html, r"(20\d{2}[-/年]\d{1,2}[-/月]\d{1,2}(?:\s+\d{1,2}:\d{2}(?::\d{2})?)?)"),
    )
    return author, media, dt


def parse_cnstock(soup: BeautifulSoup, html: str) -> tuple[str, str, str, str]:
    m = re.search(r'<script id="__NEXT_DATA__" type="application/json">(.*?)</script>', html, re.S)
    if not m:
        return "", "上海证券报·中国证券网", "", generic_body(soup)
    data = json.loads(m.group(1))
    info = (((data.get("props") or {}).get("pageProps") or {}).get("data") or {}).get("textInfo") or {}
    body = BeautifulSoup(info.get("content") or "", "lxml").get_text("\n", strip=True)
    author = pick_first(info.get("author"), info.get("authorName"), first_re(body, r"记者\s+(\S+)"))
    return author, "上海证券报·中国证券网", str(info.get("publishTime") or ""), body


def parse_ithome(soup: BeautifulSoup, html: str) -> tuple[str, str, str, str]:
    author = first_re(html, r"作者[：:]\s*(?:<[^>]+>)*\s*([^<\s]{1,20})")
    dt = first_re(html, r"(20\d{2}/\d{1,2}/\d{1,2}\s+\d{1,2}:\d{2}:\d{2})")
    body = paragraph_text(soup.select_one("#paragraph"))
    return author, "IT之家", dt, body


def parse_qq(soup: BeautifulSoup, html: str) -> tuple[str, str, str, str]:
    media = meta(soup, "article:author") or "腾讯新闻"
    dt = meta(soup, "article:published_time")
    el = soup.select_one(".rich_media_content") or soup.select_one("#article-content")
    body = paragraph_text(el)
    body = re.sub(r"^问AI[\s\S]{0,80}?(?=IT之家|【)", "", body)
    return "", media, dt, body


def parse_ifeng(soup: BeautifulSoup, html: str) -> tuple[str, str, str, str]:
    dt = meta(soup, "article:published_time", "og:time ")
    el = soup.select_one("[class*='index_text_']") or soup.select_one("[class*='index_articleBox_']")
    body = paragraph_text(el)
    media = pick_first(first_re(html, r"([\u4e00-\u9fa5A-Za-z]{2,20})\s+\d{4}年\d{2}月"), meta(soup, "og:site_name"))
    return "", media, dt, body


def parse_cet(soup: BeautifulSoup, html: str) -> tuple[str, str, str, str]:
    authors = []
    if "李海楠" in html:
        authors.append("李海楠")
    m = re.search(r"记者[^\n]{0,20}?([^\s]{2,8})[、\s]+([^\s]{2,8})", html)
    author = "、".join(authors) if authors else (f"{m.group(1)}、{m.group(2)}" if m else first_re(html, r"记者\s+(\S+)"))
    dt = first_re(html, r"(20\d{2}-\d{2}-\d{2}\s+\d{2}:\d{2}:\d{2})")
    el = soup.select_one("section.article_content") or soup.select_one(".article-body")
    paras, seen = [], set()
    nodes = (el.find_all("p", recursive=False) if el else []) or (el.find_all("p") if el else [])
    for p in nodes:
        t = p.get_text(" ", strip=True)
        t = re.sub(r"^数智热点\s*", "", t)
        t = re.sub(r"\s*■+\s*中国经济时报记者[\s\S]*$", "", t).strip()
        if not t or t.startswith("值班编辑") or t in seen:
            if t.startswith("值班编辑"):
                break
            continue
        seen.add(t)
        paras.append(t)
    return author, "中国经济时报", dt, "\n".join(paras)


SITE_PARSERS = [
    (("cnstock.com",), parse_cnstock),
    (("ithome.com",), parse_ithome),
    (("news.qq.com",), parse_qq),
    (("ifeng.com",), parse_ifeng),
    (("cet.com.cn",), parse_cet),
]


def parse_one(url: str, html_path: Path) -> dict[str, str]:
    html = decode_html(html_path)
    soup = BeautifulSoup(html, "lxml")
    host = host_of(url)
    author = media = dt = body = ""
    for hosts, fn in SITE_PARSERS:
        if any(h in host for h in hosts):
            author, media, dt, body = fn(soup, html)
            break
    if len(body) < SHORT_BODY:
        body = generic_body(soup) or body
    g_author, g_media, g_dt = generic_fields(soup, html)
    author = pick_first(author, g_author)
    media = pick_first(media, g_media)
    dt = pick_first(dt, g_dt)

    if "yiche.com" in host:
        media = media or "易车"
        el = soup.select_one(".news-detail-main")
        if el:
            body = paragraph_text(el)
    if "xinhuanet.com" in host:
        media = media or "新华网"
        el = soup.select_one("#detail")
        if el:
            body = paragraph_text(el)
        author = pick_first(author, first_re(body, r"新华社记者([\u4e00-\u9fa5、]{2,20})"), first_re(html, r"新华社记者([\u4e00-\u9fa5、]{2,20})"))
    if "sohu.com" in host:
        el = soup.select_one("#mp-editor")
        if el:
            body = paragraph_text(el)
        media = meta(soup, "mediaid") or media
    if "sciencenet.cn" in host:
        el = soup.select_one("#content1")
        if el:
            body = paragraph_text(el)
            lines = [ln for ln in body.splitlines() if ln.strip()]
            for i, ln in enumerate(lines):
                if "本报讯" in ln:
                    body = "\n".join(lines[i:])
                    break
        media = media or "中国科学报"
    if "163.com" in host:
        el = soup.select_one(".post_body")
        if el:
            body = paragraph_text(el)
        if not author:
            author = first_re(body, r"作者｜\s*(\S+)")
        if not media:
            media = first_re(body, r"文｜\s*(\S+)")
    if "tsinghua.edu.cn" in host:
        el = soup.select_one(".v_news_content")
        if el:
            body = paragraph_text(el)
        src = first_re(html, r"来源：([^<\n]{2,40})")
        if src:
            media = re.split(r"\s+", src)[0]
            dm = re.search(r"(\d{1,2})-(\d{1,2})", src)
            if dm:
                dt = dt or f"{date.today().year}-{int(dm.group(1)):02d}-{int(dm.group(2)):02d}"
            au = re.search(r"[\s　]+([\u4e00-\u9fa5]{1,4})\s*$", src)
            if au:
                author = author or au.group(1)
        body = cut_after(body, ["编辑：", "编辑:"])
    if "qianlong.com" in host:
        el = soup.select_one(".contentStr")
        if el:
            body = paragraph_text(el)
    if "sina.com" in host:
        el = soup.select_one("#artibody")
        if el:
            body = paragraph_text(el)
        body = re.sub(r"\n?\(?配置\s*\n?\|\s*\n?询价\)?", "", body)

    body = cut_after(clean_text(body))
    body = dedupe_body(body)
    if author in {"网易", "腾讯网"}:
        author = ""
    return {"author": author, "media": media, "date": normalize_date(dt), "body": body}


def write_excel(rows: list[dict[str, str]], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "正文采集"
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    alt = PatternFill("solid", fgColor="F2F7FB")
    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(1, col, h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin
    widths = {"A": 6, "B": 42, "C": 16, "D": 28, "E": 48, "F": 20, "G": 80}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(rows) + 1}"
    for i, r in enumerate(rows, 1):
        values = [i, r["title"], r["author"], r["media"], r["url"], r["date"], r["body"]]
        for col, val in enumerate(values, 1):
            c = ws.cell(i + 1, col, val)
            c.alignment = wrap
            c.font = Font(name="微软雅黑", size=10)
            c.border = thin
            if i % 2 == 0:
                c.fill = alt
        ws.row_dimensions[i + 1].height = min(220, max(48, 14 + min(r["body"].count("\n") + 1, 12) * 12))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def run(params: dict) -> dict:
    items = params.get("items")
    if not items or not isinstance(items, list):
        raise FulltextError("必须提供 items（[{title,url}] 数组）")
    patches_map: dict[str, dict] = {}
    for p in params.get("patches") or []:
        if isinstance(p, dict):
            patches_map[(p.get("url") or "").strip()] = p
    export_dir = Path(params.get("export_dir") or "./workspace")

    cache = Path(tempfile.mkdtemp(prefix="news_fulltext_"))
    rows = []
    fails = []
    total = len(items)
    print(f"开始采集正文：共 {total} 条", file=sys.stderr, flush=True)
    for i, item in enumerate(items, 1):
        print(f"采集进度 {i}/{total}", file=sys.stderr, flush=True)
        title = (item.get("title") or "").strip()
        url = (item.get("url") or "").strip()
        html_path = cache / f"{i:02d}.html"
        code = fetch_html(url, html_path) if url else 0
        parsed = {"author": "", "media": "", "date": "", "body": ""}
        if html_path.exists() and html_path.stat().st_size > 200:
            try:
                parsed = parse_one(url, html_path)
            except Exception as exc:  # noqa: BLE001
                fails.append({"index": i, "url": url, "reason": f"parse: {exc}"})
        else:
            fails.append({"index": i, "url": url, "reason": f"http={code}"})
        patch = patches_map.get(url) or {}
        row = {
            "title": title,
            "url": url,
            "author": (patch.get("author") if "author" in patch else parsed["author"]) or "",
            "media": (patch.get("media") if "media" in patch else parsed["media"]) or "",
            "date": normalize_date(patch.get("date") if "date" in patch else parsed["date"]),
            "body": (patch.get("body") if "body" in patch else parsed["body"]) or "",
        }
        rows.append(row)

    export_dir.mkdir(parents=True, exist_ok=True)
    stamp = date.today().strftime("%Y%m%d")
    out_name = params.get("filename") or f"新闻正文_{stamp}.xlsx"
    out_path = (export_dir / out_name).resolve()
    write_excel(rows, out_path)

    short_urls = [r["url"] for r in rows if len(r["body"]) < SHORT_BODY]
    return {
        "output_path": str(out_path),
        "rows": len(rows),
        "short": len(short_urls),
        "empty_author": sum(1 for r in rows if not r["author"]),
        "short_urls": short_urls,
        "failed": fails,
    }


def main() -> None:
    try:
        args = json.loads(sys.argv[1]) if len(sys.argv) > 1 and sys.argv[1] else {}
        data = run(args if isinstance(args, dict) else {})
        print(json.dumps({"ok": True, "data": data}, ensure_ascii=False))
    except FulltextError as err:
        print(json.dumps({"ok": False, "error": str(err), "hint": "提供 items=[{title,url}]"}, ensure_ascii=False))
    except Exception as err:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": f"{type(err).__name__}: {err}"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
