#!/usr/bin/env python3
"""SKE 社媒打标：读取 Excel/CSV 指定列，按提示词逐行生成列，同时导出 xlsx + csv。

默认提示词为 SKE 社媒字段抽取（media/date/author/sentiment/summary/分类/原链接）。
也可传入自定义 prompt + output_columns。密钥仅从环境变量读取。
"""

from __future__ import annotations

import csv
import json
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook

DEFAULT_OUTPUT_COLUMNS = ["media", "date", "author", "sentiment", "summary", "分类", "原链接"]
KNOWN_MEDIA = ("Instagram", "Facebook", "X(Twitter)", "Tumblr", "Youtube", "Reddit", "Bluesky")
MAX_RETRIES = 2
RETRY_BACKOFF_SECONDS = 1.5
TEXT_MAX_CHARS = 8000

DEFAULT_PROMPT = """media：从<文本>列中识别所带的第一个超链接，并且将其主域名分类为来源，去除".com"，首字母大写。如果是“t.co”或“twitter.com”则识别为“X(Twitter)”；如果是“reddit.com”则为Reddit。类别包含：Instagram、Facebook、X(Twitter)、Tumblr、Youtube、Reddit、Bluesky。如果第一个超链接无法识别出主域名的情况标记为Reddit。

date:提取<文本>中的第一个日期，显示为 YYYY-MM-DD格式。年份为当前年。

author: <文本>仅提取作者名称（如有@符号，取@之后内容），不要来源、正文或摘要。如果<media>中的文本不是“Instagram、Facebook、X(Twitter)、Tumblr、Youtube、Reddit、Bluesky”的其中之一，并且未在 <文本> 提取出作者名，则标记为：“Forum User”。

sentiment:<文本>从文本中提取情感分类，分类包括：positive/negative/neutral。

summary: 查看<文本>仅提取正文内容或摘要，只保留包含以下关键词之一的完整语句：SKE,Hayati,Vaporesso,RELX,Geek Bar,Elfbar
注意：不包含来源、日期、作者等信息。
将所有非英语的语言翻译成英文。

分类： 查看<文本>内容是否命中以下关键字，提取主要的一个关键字作为分类，如未命中则置为空：
SKE,Hayati,Vaporesso,RELX,Geek Bar,Elfbar

原链接：从 <文本> 字段中提取所有链接或域名。

规则：
1. 只输出链接本身，不要解释。
2. 如果原文中有 http:// 或 https:// 开头的链接，必须完整保留原始链接，包括路径、参数、#、短链后缀，不要省略为“...”，不要改写，不要展开短链。
3. 如果原文中没有 http:// 或 https://，但有域名形式的内容，例如 example.com、abc.co.uk、xxx.tumblr.com，也提取该域名。
4. 如果同一条文本中有多个链接，用英文逗号加空格分隔。
5. 链接结尾不要包含多余的中文/英文标点、括号、引号、空格。
6. 如果没有任何链接或域名，输出“无”。

输出格式示例：
https://t.co/abc123, example.com"""


class TaggingError(Exception):
    pass


def get_runtime_config() -> dict[str, str]:
    api_key = (os.getenv("DEEPSEEK_API_KEY") or "").strip()
    if not api_key:
        raise TaggingError("缺少 DEEPSEEK_API_KEY 环境变量")
    base_url = (os.getenv("DEEPSEEK_BASE_URL") or "https://api.deepseek.com").rstrip("/")
    model = os.getenv("DEEPSEEK_MODEL") or "deepseek-v4-flash"
    return {"api_key": api_key, "base_url": base_url, "model": model}


def resolve_input(path_str: str) -> Path:
    raw = Path(str(path_str)).expanduser()
    if raw.is_file():
        return raw.resolve()
    cwd = Path.cwd()
    name = Path(str(path_str)).name
    for cand in (cwd / str(path_str), cwd / "uploads" / name, cwd / "workspace" / name):
        if cand.is_file():
            return cand.resolve()
    raise TaggingError(f"输入文件不存在: {path_str}")


def empty_result(columns: list[str]) -> dict[str, str]:
    out = {col: "" for col in columns}
    if "原链接" in out:
        out["原链接"] = "无"
    return out


def normalize_result(raw: dict[str, Any], columns: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for col in columns:
        value = raw.get(col, "")
        result[col] = "" if value is None else str(value).strip()
    if "sentiment" in result:
        sent = result["sentiment"].lower()
        result["sentiment"] = sent if sent in {"positive", "negative", "neutral"} else "neutral"
    if "media" in result and result["media"]:
        media = result["media"].strip()
        aliases = {
            "twitter": "X(Twitter)",
            "x": "X(Twitter)",
            "x(twitter)": "X(Twitter)",
            "youtube": "Youtube",
            "you tube": "Youtube",
        }
        result["media"] = aliases.get(media.lower(), media)
    if "原链接" in result and not result["原链接"]:
        result["原链接"] = "无"
    if "author" in result:
        media = result.get("media", "")
        if not result["author"] and media not in KNOWN_MEDIA:
            result["author"] = "Forum User"
    return result


def extract_json_object(text: str) -> dict[str, Any]:
    stripped = (text or "").strip()
    if stripped.startswith("```"):
        stripped = re.sub(r"^```(?:json)?\s*", "", stripped, flags=re.IGNORECASE)
        stripped = re.sub(r"\s*```$", "", stripped)
    start = stripped.find("{")
    end = stripped.rfind("}")
    if start == -1 or end == -1 or end <= start:
        raise ValueError("模型未返回 JSON 对象")
    parsed = json.loads(stripped[start : end + 1])
    if not isinstance(parsed, dict):
        raise ValueError("模型返回的 JSON 不是对象")
    return parsed


def build_user_prompt(task_prompt: str, text_col: str, text: str, columns: list[str], year: int) -> str:
    cols = json.dumps(columns, ensure_ascii=False)
    clipped = text if len(text) <= TEXT_MAX_CHARS else text[:TEXT_MAX_CHARS] + "\n…(truncated)"
    return f"""你是表格字段抽取助手。请严格按【提示词】从指定列抽取/生成字段。

当前年份：{year}
指定列名：{text_col}

【提示词】
{task_prompt}

【{text_col}】
{clipped}

【输出要求】
只输出一个 JSON 对象，键必须恰好为：{cols}
所有值都是字符串。不要 markdown、不要解释、不要额外键。""".strip()


def call_llm(cfg: dict[str, str], user_prompt: str, columns: list[str]) -> dict[str, str]:
    url = f"{cfg['base_url']}/chat/completions"
    body = {
        "model": cfg["model"],
        "messages": [{"role": "user", "content": user_prompt}],
        "response_format": {"type": "json_object"},
        "extra_body": {"thinking": {"type": "disabled"}},
    }
    headers = {"Content-Type": "application/json", "Authorization": f"Bearer {cfg['api_key']}"}
    last_error = ""
    for attempt in range(MAX_RETRIES + 1):
        try:
            req = Request(
                url,
                method="POST",
                headers=headers,
                data=json.dumps(body, ensure_ascii=False).encode("utf-8"),
            )
            with urlopen(req, timeout=90) as resp:
                data = json.loads(resp.read().decode("utf-8", errors="ignore"))
            content = data["choices"][0]["message"]["content"]
            return normalize_result(extract_json_object(content), columns)
        except HTTPError as err:
            detail = err.read().decode("utf-8", errors="ignore")
            last_error = f"HTTP {err.code}: {detail[:500]}"
            if err.code in (429, 500, 502, 503, 504) and attempt < MAX_RETRIES:
                time.sleep(RETRY_BACKOFF_SECONDS * (2**attempt))
                continue
        except Exception as err:  # noqa: BLE001
            last_error = str(err)
            if attempt < MAX_RETRIES:
                time.sleep(RETRY_BACKOFF_SECONDS * (2**attempt))
                continue
        break
    failed = empty_result(columns)
    if "sentiment" in failed:
        failed["sentiment"] = "neutral"
    if "summary" in failed:
        failed["summary"] = f"错误: {last_error}"[:400]
    return failed


def parse_output_columns(params: dict) -> list[str]:
    raw = params.get("output_columns")
    if raw is None or raw == "":
        return list(DEFAULT_OUTPUT_COLUMNS)
    if isinstance(raw, str):
        parts = [p.strip() for p in re.split(r"[,，]", raw) if p.strip()]
        if not parts:
            raise TaggingError("output_columns 不能为空")
        return parts
    if isinstance(raw, list):
        cols = [str(x).strip() for x in raw if str(x).strip()]
        if not cols:
            raise TaggingError("output_columns 不能为空")
        return cols
    raise TaggingError("output_columns 必须是字符串数组或逗号分隔字符串")


def read_table(src: Path, sheet: str | None) -> tuple[list[str], list[dict[str, str]]]:
    ext = src.suffix.lower()
    if ext == ".csv":
        with src.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            if not reader.fieldnames:
                raise TaggingError("CSV 未检测到表头")
            headers = [str(h) for h in reader.fieldnames]
            rows: list[dict[str, str]] = []
            for row in reader:
                rows.append({h: "" if row.get(h) is None else str(row.get(h)) for h in headers})
            return headers, rows
    if ext in {".xlsx", ".xlsm"}:
        wb = load_workbook(src, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        raw_headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
        headers = []
        used: dict[str, int] = {}
        for idx, value in enumerate(raw_headers, start=1):
            base = ("" if value is None else str(value)).strip() or f"col_{idx}"
            n = used.get(base, 0) + 1
            used[base] = n
            headers.append(base if n == 1 else f"{base}_{n}")
        rows = []
        for row_idx in range(2, ws.max_row + 1):
            item: dict[str, str] = {}
            empty = True
            for col_idx, name in enumerate(headers, start=1):
                val = ws.cell(row=row_idx, column=col_idx).value
                text = "" if val is None else str(val)
                item[name] = text
                if text.strip():
                    empty = False
            if not empty:
                rows.append(item)
        return headers, rows
    raise TaggingError(f"仅支持 .xlsx / .csv，收到: {ext or '(无扩展名)'}")


def write_xlsx(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "tagged"
    for col_idx, name in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=name)
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, name in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row.get(name, ""))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_csv(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({h: row.get(h, "") for h in headers})


def run(params: dict) -> dict:
    input_path = params.get("input_path") or params.get("input") or params.get("source_path")
    if not input_path:
        raise TaggingError("必须提供 input_path（待处理的 Excel 或 CSV）")
    src = resolve_input(str(input_path))
    sheet = params.get("sheet") or params.get("sheet_name")
    text_col = str(params.get("text_col") or params.get("column") or "文本").strip() or "文本"
    task_prompt = str(params.get("prompt") or DEFAULT_PROMPT).strip() or DEFAULT_PROMPT
    columns = parse_output_columns(params)
    max_workers = max(1, min(16, int(params.get("max_workers", 8))))
    max_rows = int(params.get("max_rows") or 0)
    export_dir = Path(params.get("export_dir") or "./workspace")
    year = int(params.get("year") or datetime.now().year)

    headers, rows = read_table(src, str(sheet) if sheet else None)
    if text_col not in headers:
        if len(headers) == 1:
            text_col = headers[0]
        else:
            raise TaggingError(f"找不到指定列「{text_col}」，当前表头={headers}")
    if max_rows > 0:
        rows = rows[:max_rows]

    out_headers = list(headers)
    for col in columns:
        if col not in out_headers:
            out_headers.append(col)

    cfg = get_runtime_config()
    total = len(rows)
    print(f"开始 SKE 社媒打标：共 {total} 行，并发 {max_workers}", file=sys.stderr, flush=True)

    indexed: list[tuple[int, str]] = []
    for i, row in enumerate(rows):
        text = (row.get(text_col) or "").strip()
        indexed.append((i, text))

    results: dict[int, dict[str, str]] = {}
    done = 0

    def _one(idx: int, text: str) -> dict[str, str]:
        if not text:
            return empty_result(columns)
        prompt = build_user_prompt(task_prompt, text_col, text, columns, year)
        return call_llm(cfg, prompt, columns)

    with ThreadPoolExecutor(max_workers=max_workers) as pool:
        fut_map = {pool.submit(_one, idx, text): idx for idx, text in indexed}
        for fut in as_completed(fut_map):
            idx = fut_map[fut]
            try:
                results[idx] = fut.result()
            except Exception as err:  # noqa: BLE001
                failed = empty_result(columns)
                if "summary" in failed:
                    failed["summary"] = f"错误: {err}"[:400]
                results[idx] = failed
            done += 1
            if done == total or done % 5 == 0:
                print(f"打标进度 {done}/{total}", file=sys.stderr, flush=True)

    errors = 0
    for idx, row in enumerate(rows):
        labels = results.get(idx) or empty_result(columns)
        if str(labels.get("summary") or "").startswith("错误"):
            errors += 1
        for col in columns:
            row[col] = labels.get(col, "")

    export_dir.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", src.stem).strip("._") or "ske_tagged"
    xlsx_path = (export_dir / f"{stem}_ske_tagged.xlsx").resolve()
    csv_path = (export_dir / f"{stem}_ske_tagged.csv").resolve()
    write_xlsx(xlsx_path, out_headers, rows)
    write_csv(csv_path, out_headers, rows)

    return {
        "output_xlsx": str(xlsx_path),
        "output_csv": str(csv_path),
        "rows": total,
        "errors": errors,
        "text_col": text_col,
        "columns_added": columns,
        "note": "请点击界面下载按钮获取 xlsx / csv，不要编造下载链接。",
    }


def main() -> None:
    try:
        raw = sys.argv[1] if len(sys.argv) > 1 else ""
        args = json.loads(raw) if raw.strip() else {}
        if not isinstance(args, dict):
            raise TaggingError("参数 JSON 必须是对象")
        data = run(args)
        print(json.dumps({"ok": True, "data": data}, ensure_ascii=False))
    except TaggingError as err:
        print(json.dumps(
            {
                "ok": False,
                "error": str(err),
                "hint": "确认已上传 xlsx/csv、指定列存在，以及 DEEPSEEK_API_KEY 已配置",
            },
            ensure_ascii=False,
        ))
    except Exception as err:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": f"{type(err).__name__}: {err}"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
