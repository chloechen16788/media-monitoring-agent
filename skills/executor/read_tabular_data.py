"""读取结构化表格数据（xlsx/csv/jsonl）并标准化为行对象。"""

import csv
import json
import os
import sys
from typing import Any

from openpyxl import load_workbook


def _emit(payload: dict[str, Any]) -> None:
    print(json.dumps(payload, ensure_ascii=False))


def fail(error: str, hint: str = "") -> None:
    payload: dict[str, Any] = {"ok": False, "error": error}
    if hint:
        payload["hint"] = hint
    _emit(payload)
    sys.exit(0)


def read_params() -> dict[str, Any]:
    raw = sys.argv[1] if len(sys.argv) > 1 else ""
    if not raw.strip():
        return {}
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        fail("参数必须是单个合法的 JSON 字符串（禁止单引号）。")
    if not isinstance(parsed, dict):
        fail("参数 JSON 必须是对象。")
    return parsed


def _safe_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _sanitize_headers(raw_headers: list[Any]) -> list[str]:
    headers: list[str] = []
    used: dict[str, int] = {}
    for idx, value in enumerate(raw_headers, start=1):
        base = _safe_str(value).strip() or f"col_{idx}"
        count = used.get(base, 0) + 1
        used[base] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return headers


def _pick_fields(row: dict[str, Any], fields: list[str] | None) -> dict[str, Any]:
    if not fields:
        return row
    return {k: row.get(k, "") for k in fields}


def _parse_max_rows(params: dict[str, Any]) -> int:
    raw = params.get("max_rows")
    if raw is None or raw == "":
        return 0
    try:
        value = int(raw)
    except (TypeError, ValueError):
        fail("max_rows 必须是非负整数。")
    if value < 0:
        fail("max_rows 必须是非负整数。")
    return value


def read_xlsx(params: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    source_path = str(params["source_path"])
    sheet_name = params.get("sheet_name")
    sheet_index = int(params.get("sheet_index", 0))
    header_row = int(params.get("header_row", 1))
    start_row = int(params.get("start_row", header_row + 1))
    max_rows = _parse_max_rows(params)
    fields = params.get("fields")
    if fields is not None and not isinstance(fields, list):
        fail("fields 必须是字符串数组。")
    if isinstance(fields, list) and any(not isinstance(x, str) for x in fields):
        fail("fields 必须是字符串数组。")

    wb = load_workbook(source_path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            fail(f"找不到工作表: {sheet_name}")
        ws = wb[sheet_name]
    else:
        names = wb.sheetnames
        if not names:
            fail("xlsx 文件不包含任何工作表。")
        if sheet_index < 0 or sheet_index >= len(names):
            fail(f"sheet_index 越界: {sheet_index}")
        ws = wb[names[sheet_index]]

    rows_iter = ws.iter_rows(values_only=True)
    header_values: list[Any] | None = None
    for idx, values in enumerate(rows_iter, start=1):
        if idx == header_row:
            header_values = list(values)
            break
    if header_values is None:
        fail("无法读取表头行，请检查 header_row。")

    headers = _sanitize_headers(header_values)
    rows: list[dict[str, Any]] = []
    logical_row_id = 1
    truncated = False

    for idx, values in enumerate(ws.iter_rows(min_row=start_row, values_only=True), start=start_row):
        row_obj = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
        row_obj = _pick_fields(row_obj, fields)
        row_obj["__row_id"] = logical_row_id
        row_obj["__source_row_number"] = idx
        rows.append(row_obj)
        logical_row_id += 1
        if max_rows > 0 and len(rows) >= max_rows:
            truncated = True
            break

    return rows, {
        "file_type": "xlsx",
        "sheet_name": ws.title,
        "header_row": header_row,
        "start_row": start_row,
        "row_count": len(rows),
        "columns": headers if not fields else fields,
        "max_rows_applied": max_rows,
        "truncated": truncated,
    }


def read_csv_file(params: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    source_path = str(params["source_path"])
    encoding = str(params.get("encoding", "utf-8-sig"))
    delimiter = str(params.get("delimiter", ","))
    max_rows = _parse_max_rows(params)
    fields = params.get("fields")
    if fields is not None and not isinstance(fields, list):
        fail("fields 必须是字符串数组。")
    if isinstance(fields, list) and any(not isinstance(x, str) for x in fields):
        fail("fields 必须是字符串数组。")

    rows: list[dict[str, Any]] = []
    truncated = False
    with open(source_path, "r", encoding=encoding, newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        if not reader.fieldnames:
            fail("CSV 未检测到表头。")
        for idx, row in enumerate(reader, start=1):
            norm_row = {k: row.get(k, "") for k in reader.fieldnames}
            norm_row = _pick_fields(norm_row, fields)
            norm_row["__row_id"] = idx
            norm_row["__source_row_number"] = idx + 1
            rows.append(norm_row)
            if max_rows > 0 and len(rows) >= max_rows:
                truncated = True
                break

    columns = list(reader.fieldnames or [])
    return rows, {
        "file_type": "csv",
        "row_count": len(rows),
        "columns": columns if not fields else fields,
        "encoding": encoding,
        "max_rows_applied": max_rows,
        "truncated": truncated,
    }


def read_jsonl_file(params: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    source_path = str(params["source_path"])
    max_rows = _parse_max_rows(params)
    fields = params.get("fields")
    if fields is not None and not isinstance(fields, list):
        fail("fields 必须是字符串数组。")
    if isinstance(fields, list) and any(not isinstance(x, str) for x in fields):
        fail("fields 必须是字符串数组。")

    rows: list[dict[str, Any]] = []
    truncated = False
    with open(source_path, "r", encoding="utf-8") as f:
        for source_line_no, line in enumerate(f, start=1):
            text = line.strip()
            if not text:
                continue
            try:
                obj = json.loads(text)
            except json.JSONDecodeError:
                continue
            if not isinstance(obj, dict):
                continue
            row = _pick_fields(dict(obj), fields)
            row["__row_id"] = len(rows) + 1
            row["__source_row_number"] = source_line_no
            rows.append(row)
            if max_rows > 0 and len(rows) >= max_rows:
                truncated = True
                break

    columns = sorted({k for row in rows for k in row.keys() if not k.startswith("__")})
    return rows, {
        "file_type": "jsonl",
        "row_count": len(rows),
        "columns": columns if not fields else fields,
        "max_rows_applied": max_rows,
        "truncated": truncated,
    }


def infer_file_type(path: str, explicit: str | None) -> str:
    if explicit:
        return explicit.lower()
    ext = os.path.splitext(path)[1].lower()
    if ext in [".xlsx", ".xlsm"]:
        return "xlsx"
    if ext == ".csv":
        return "csv"
    if ext == ".jsonl":
        return "jsonl"
    fail("无法自动识别文件类型。", "请传 file_type: xlsx/csv/jsonl。")
    return ""


def main() -> None:
    params = read_params()
    source_path = str(params.get("source_path", "")).strip()
    if not source_path:
        fail("缺少必填参数: source_path")
    if not os.path.exists(source_path):
        fail(f"文件不存在: {source_path}")

    file_type = infer_file_type(source_path, params.get("file_type"))
    if file_type == "xlsx":
        rows, meta = read_xlsx(params)
    elif file_type == "csv":
        rows, meta = read_csv_file(params)
    elif file_type == "jsonl":
        rows, meta = read_jsonl_file(params)
    else:
        fail(f"不支持的 file_type: {file_type}", "支持 xlsx/csv/jsonl。")
        return

    _emit(
        {
            "ok": True,
            "data": {
                "source_path": os.path.abspath(source_path),
                "rows": rows,
                "meta": meta,
            },
        }
    )


if __name__ == "__main__":
    main()
