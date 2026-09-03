#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fill Xiaohongshu engagement metrics from note links in an uploaded table.

Reads Excel/CSV, extracts note_id from the link column, calls TikHub
get_image_note_detail / get_video_note_detail, writes metrics back onto a copy
of the original table. Does not call other skills.
"""

from __future__ import annotations

import csv
import json
import re
import sys
import threading
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from _xiaohongshu_tikhub import (
    DEFAULT_QPS,
    DEFAULT_WORKERS,
    DETAIL_IMAGE_PATH,
    DETAIL_VIDEO_PATH,
    MAX_WORKERS,
    USER_INFO_PATH,
    TikHub,
    XhsError,
    build_cost_from_paths,
    detail_path,
    extract_detail_metrics,
    extract_note_id_from_text,
    extract_urls,
    ensure_xlsx_filename,
    fill_fans,
    guess_note_kind,
    is_fatal_tikhub,
    media_type,
    optional_env,
    proxy_host,
    require_env,
    uniq_warnings,
)

LINK_COL_ALIASES = (
    "笔记链接",
    "笔记地址",
    "原链接",
    "链接",
    "地址",
    "url",
    "URL",
    "note_url",
    "note_link",
    "link",
)
METRIC_SPECS = (
    ("note_id", ("note_id", "笔记ID", "笔记id")),
    ("liked", ("点赞数", "点赞")),
    ("collects", ("收藏数", "收藏")),
    ("comments", ("评论数", "评论")),
    ("shares", ("分享数", "转发数", "转发")),
    ("media", ("类型", "媒体类型")),
    ("status", ("详情状态",)),
    ("error", ("详情错误",)),
)
OPTIONAL_FILL = (
    ("title", ("标题",)),
    ("nickname", ("作者号", "作者", "昵称")),
    ("time", ("时间", "发布时间")),
    ("body", ("正文",)),
    ("fans", ("粉丝数",)),
)


class LinkDetailError(XhsError):
    """User-facing failure for the link-detail skill."""


def resolve_input(path_str: str) -> Path:
    raw = Path(str(path_str)).expanduser()
    if raw.is_file():
        return raw.resolve()
    cwd = Path.cwd()
    name = Path(str(path_str)).name
    for cand in (cwd / str(path_str), cwd / "uploads" / name, cwd / "workspace" / name):
        if cand.is_file():
            return cand.resolve()
    raise LinkDetailError(f"输入文件不存在: {path_str}")


def cell_text(cell: Any) -> str:
    hyper = getattr(cell, "hyperlink", None)
    target = getattr(hyper, "target", None) if hyper is not None else None
    if target:
        return str(target).strip()
    value = cell.value
    return "" if value is None else str(value).strip()


def sanitize_headers(raw_headers: list[Any]) -> list[str]:
    headers: list[str] = []
    used: dict[str, int] = {}
    for idx, value in enumerate(raw_headers, start=1):
        base = ("" if value is None else str(value)).strip() or f"col_{idx}"
        count = used.get(base, 0) + 1
        used[base] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return headers


def read_table(src: Path, sheet: str | None) -> tuple[list[str], list[dict[str, str]]]:
    ext = src.suffix.lower()
    if ext == ".csv":
        with src.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            if not reader.fieldnames:
                raise LinkDetailError("CSV 未检测到表头")
            headers = [str(h) for h in reader.fieldnames]
            rows: list[dict[str, str]] = []
            for row in reader:
                item = {h: "" if row.get(h) is None else str(row.get(h)) for h in headers}
                if any(str(v).strip() for v in item.values()):
                    rows.append(item)
            return headers, rows
    if ext in {".xlsx", ".xlsm"}:
        wb = load_workbook(src, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        raw_headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
        headers = sanitize_headers(raw_headers)
        rows = []
        for row_idx in range(2, ws.max_row + 1):
            item = {}
            empty = True
            for col_idx, name in enumerate(headers, start=1):
                text = cell_text(ws.cell(row=row_idx, column=col_idx))
                item[name] = text
                if text:
                    empty = False
            if not empty:
                rows.append(item)
        return headers, rows
    raise LinkDetailError(f"仅支持 .xlsx / .csv，收到: {ext or '(无扩展名)'}")


def pick_link_col(headers: list[str], requested: str) -> str:
    if requested and requested in headers:
        return requested
    if requested:
        lowered = {h.lower(): h for h in headers}
        if requested.lower() in lowered:
            return lowered[requested.lower()]
        raise LinkDetailError(f"找不到指定列「{requested}」，当前表头={headers}")
    for alias in LINK_COL_ALIASES:
        if alias in headers:
            return alias
    lowered = {h.lower(): h for h in headers}
    for alias in LINK_COL_ALIASES:
        if alias.lower() in lowered:
            return lowered[alias.lower()]
    raise LinkDetailError(
        f"未找到笔记链接列，请传 link_col。当前表头={headers}"
    )


def pick_existing(headers: list[str], aliases: tuple[str, ...]) -> str | None:
    for name in aliases:
        if name in headers:
            return name
    lowered = {h.lower(): h for h in headers}
    for name in aliases:
        if name.lower() in lowered:
            return lowered[name.lower()]
    return None


def share_text_for(cell: str, note_id: str) -> str:
    urls = extract_urls(cell)
    for url in urls:
        if "xiaohongshu.com" in url.lower() or "xhslink.com" in url.lower():
            return url
    if urls:
        return urls[0]
    if "xhslink.com" in (cell or "").lower() or "xiaohongshu.com" in (cell or "").lower():
        return cell.strip()
    return note_id


def fatal_tikhub(err: Exception) -> bool:
    return is_fatal_tikhub(err)


def fetch_one_detail(
    api: TikHub, note_id: str, share_text: str, preferred: str
) -> tuple[dict[str, Any], str, str]:
    params: dict[str, Any] = {}
    if note_id:
        params["note_id"] = note_id
    if share_text and share_text != note_id:
        params["share_text"] = share_text
    if not params:
        return {}, "", "empty"
    order = ["video", "image"] if preferred == "video" else ["image", "video"]
    last_err = ""
    for kind in order:
        try:
            data = api.get(detail_path(kind), params)
        except XhsError as err:
            if fatal_tikhub(err):
                raise
            last_err = str(err)
            continue
        if data.get("code") != 200:
            msg = data.get("message_zh") or data.get("message") or data.get("msg") or ""
            last_err = f"TikHub {data.get('code')}: {msg}".strip()
            continue
        metrics = extract_detail_metrics(data)
        if not metrics.get("note_id"):
            metrics["note_id"] = note_id
        if kind == "video" and not str(metrics.get("type") or "").strip():
            metrics["type"] = "video"
        return metrics, kind, ""
    return {}, "", last_err or "详情失败"


def write_table(path: Path, headers: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({h: row.get(h, "") for h in headers})
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "互动数据"
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.auto_filter.ref = f"A1:{chr(ord('A') + min(len(headers), 26) - 1)}{ws.max_row}"
    ws.freeze_panes = "A2"
    wb.save(path)


def default_filename(stem: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    slug = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", stem).strip("._") or "note_detail"
    return f"{slug}_互动数据-{stamp}.xlsx"


def run(params: dict[str, Any]) -> dict[str, Any]:
    input_path = params.get("input_path") or params.get("source_path") or params.get("input")
    if not input_path:
        raise LinkDetailError("必须提供 input_path（含笔记链接列的 Excel/CSV）")
    src = resolve_input(str(input_path))
    sheet = params.get("sheet") or params.get("sheet_name")
    requested_col = str(params.get("link_col") or params.get("text_col") or "").strip()
    with_body = bool(params.get("with_body"))
    with_fans = bool(params.get("with_fans"))
    try:
        workers = int(params.get("workers") or DEFAULT_WORKERS)
    except (TypeError, ValueError) as err:
        raise LinkDetailError("workers 必须是整数") from err
    workers = max(1, min(workers, MAX_WORKERS, DEFAULT_QPS))
    try:
        max_rows = int(params.get("max_rows") or 0)
    except (TypeError, ValueError) as err:
        raise LinkDetailError("max_rows 必须是整数") from err
    export_dir = Path(params.get("export_dir") or "./workspace")
    filename = ensure_xlsx_filename(
        str(params.get("filename") or "").strip() or default_filename(src.stem)
    )
    if "/" in filename or "\\" in filename:
        raise LinkDetailError("filename 只能是文件名，不能含路径")

    headers, rows = read_table(src, str(sheet) if sheet else None)
    if not rows:
        raise LinkDetailError("表格没有数据行")
    link_col = pick_link_col(headers, requested_col)
    if max_rows > 0:
        rows = rows[:max_rows]

    out_headers = list(headers)
    col_map: dict[str, str] = {}
    extra_specs = list(METRIC_SPECS)
    if with_body:
        extra_specs.append(("body", ("正文",)))
    if with_fans:
        extra_specs.append(("fans", ("粉丝数",)))
    for key, aliases in extra_specs:
        existing = pick_existing(out_headers, aliases)
        name = existing or aliases[0]
        col_map[key] = name
        if name not in out_headers:
            out_headers.append(name)

    jobs: list[dict[str, Any]] = []
    seen: OrderedDict[str, dict[str, Any]] = OrderedDict()
    type_col = pick_existing(headers, ("类型", "媒体类型"))
    for idx, row in enumerate(rows):
        cell = (row.get(link_col) or "").strip()
        note_id = extract_note_id_from_text(cell) or extract_note_id_from_text(row.get("笔记id") or row.get("笔记ID") or row.get("note_id") or "")
        share = share_text_for(cell, note_id)
        kind = guess_note_kind(cell, row.get(type_col) or "" if type_col else "")
        if not note_id and "xhslink.com" not in cell.lower() and "xiaohongshu.com" not in cell.lower():
            jobs.append({"idx": idx, "key": "", "note_id": "", "share": "", "kind": kind, "skip": "invalid_link"})
            continue
        key = note_id or share
        if key not in seen:
            seen[key] = {"note_id": note_id, "share": share, "kind": kind}
        jobs.append({"idx": idx, "key": key, "note_id": note_id, "share": share, "kind": kind, "skip": ""})

    unique_jobs = list(seen.items())
    print(
        f"rows={len(rows)} unique={len(unique_jobs)} link_col={link_col} "
        f"with_body={with_body} with_fans={with_fans} workers={workers}",
        file=sys.stderr,
        flush=True,
    )

    results: dict[str, dict[str, Any]] = {}
    warnings: list[str] = []
    detail_ok = 0
    detail_fail = 0
    billed_api: TikHub | None = None
    if unique_jobs:
        proxy = optional_env("SKILL_TIKHUB_PROXY")
        api = TikHub(
            require_env("SKILL_TIKHUB_API_KEY"),
            require_env("SKILL_TIKHUB_BASE_URL"),
            DEFAULT_QPS,
            proxy=proxy,
        )
        billed_api = api
        stop = threading.Event()
        print("proxy=" + (proxy_host(proxy) if proxy else "off"), file=sys.stderr, flush=True)
        try:
            def one(item: tuple[str, dict[str, Any]]) -> tuple[str, dict[str, Any], str]:
                key, job = item
                if stop.is_set():
                    return key, {}, "skipped: TikHub 中断"
                metrics, _used, err = fetch_one_detail(api, job["note_id"], job["share"], job["kind"])
                return key, metrics, err

            with ThreadPoolExecutor(max_workers=workers) as pool:
                futs = {pool.submit(one, item): item[0] for item in unique_jobs}
                done = 0
                for fut in as_completed(futs):
                    key = futs[fut]
                    done += 1
                    try:
                        _, metrics, err = fut.result()
                    except Exception as err:  # noqa: BLE001
                        err_text = str(err)
                        if fatal_tikhub(err):
                            stop.set()
                        results[key] = {"status": "fail", "error": err_text}
                        detail_fail += 1
                        warnings.append(err_text)
                        print(f"DETAIL_FAIL {key} {err}", file=sys.stderr, flush=True)
                        continue
                    if err or not metrics:
                        results[key] = {"status": "fail", "error": err or "详情为空"}
                        detail_fail += 1
                        if err:
                            warnings.append(err)
                    else:
                        if with_body:
                            metrics["body"] = metrics.get("desc") or metrics.get("title") or ""
                        metrics["status"] = "ok"
                        metrics["error"] = ""
                        results[key] = metrics
                        detail_ok += 1
                    if done % 20 == 0 or done == len(unique_jobs):
                        print(
                            f"detail {done}/{len(unique_jobs)} ok={detail_ok} fail={detail_fail}",
                            file=sys.stderr,
                            flush=True,
                        )

            if with_fans:
                rec_map: OrderedDict[str, dict[str, Any]] = OrderedDict()
                for key, metrics in results.items():
                    if metrics.get("status") != "ok":
                        continue
                    rec_map[key] = {
                        "user_id": metrics.get("user_id") or "",
                        "note_id": metrics.get("note_id") or key,
                    }
                fans_map, fan_warns = fill_fans(api, rec_map, workers=workers)
                warnings.extend(fan_warns)
                for key, metrics in results.items():
                    uid = str(metrics.get("user_id") or "")
                    metrics["fans"] = fans_map.get(uid, 0) if uid else 0
        finally:
            api.close()

    filled_rows: list[dict[str, str]] = []
    invalid = 0
    for job in jobs:
        row = dict(rows[job["idx"]])
        if job["skip"] == "invalid_link":
            invalid += 1
            payload = {"status": "invalid_link", "error": "无法从该单元格解析笔记链接或 note_id"}
        else:
            payload = results.get(job["key"]) or {"status": "fail", "error": "未取到详情"}
        row[col_map["note_id"]] = str(payload.get("note_id") or job["note_id"] or "")
        row[col_map["liked"]] = "" if payload.get("status") != "ok" else str(payload.get("liked", 0))
        row[col_map["collects"]] = "" if payload.get("status") != "ok" else str(payload.get("collects", 0))
        row[col_map["comments"]] = "" if payload.get("status") != "ok" else str(payload.get("comments", 0))
        row[col_map["shares"]] = "" if payload.get("status") != "ok" else str(payload.get("shares", 0))
        row[col_map["media"]] = media_type(str(payload.get("type") or job["kind"])) if payload.get("status") == "ok" else ""
        row[col_map["status"]] = str(payload.get("status") or "fail")
        row[col_map["error"]] = str(payload.get("error") or "")
        if with_body and "body" in col_map:
            row[col_map["body"]] = str(payload.get("body") or "")
        if with_fans and "fans" in col_map:
            row[col_map["fans"]] = "" if payload.get("status") != "ok" else str(payload.get("fans", 0) or 0)
        for field, aliases in OPTIONAL_FILL:
            if field in {"body", "fans"}:
                continue
            existing = pick_existing(headers, aliases)
            if existing and not str(row.get(existing) or "").strip() and payload.get("status") == "ok":
                value = payload.get(field)
                if value:
                    row[existing] = str(value)
        filled_rows.append(row)

    billed = billed_api.billed if billed_api is not None else None
    cost = build_cost_from_paths(
        billed or {},
        [
            (DETAIL_IMAGE_PATH, "图文详情 get_image_note_detail"),
            (DETAIL_VIDEO_PATH, "视频详情 get_video_note_detail"),
            (USER_INFO_PATH, "get_user_info 拉粉丝"),
        ],
    )
    warn_list = uniq_warnings(warnings)
    export_dir.mkdir(parents=True, exist_ok=True)
    out_path = (export_dir / filename).resolve()
    write_table(out_path, out_headers, filled_rows)

    produced = [out_path.name]
    csv_name = ""
    if src.suffix.lower() == ".csv" or filename.lower().endswith(".xlsx"):
        csv_path = out_path.with_suffix(".csv")
        write_table(csv_path, out_headers, filled_rows)
        csv_name = csv_path.name
        produced.append(csv_name)

    ok_rows = sum(1 for row in filled_rows if row.get(col_map["status"]) == "ok")
    note = "请点击下载按钮获取文件。"
    if warn_list:
        note = (
            f"中途中断（常见原因：TikHub 付费余额不足），已导出 {ok_rows}/{len(filled_rows)} 条成功详情。"
            f"{'；'.join(warn_list)} 请点击下载按钮获取文件。"
        )
    note += f" 本次消耗约 ${cost['total_usd']:.2f}，回复中须贴出 data.cost.table。"
    return {
        "output_xlsx": out_path.name if out_path.suffix.lower() != ".csv" else csv_name,
        "output_csv": csv_name,
        "output_path": str(out_path),
        "rows": len(filled_rows),
        "unique_notes": len(unique_jobs),
        "detail_ok": detail_ok,
        "detail_fail": detail_fail,
        "invalid_link": invalid,
        "link_col": link_col,
        "columns_added": [name for name in out_headers if name not in headers],
        "with_body": with_body,
        "with_fans": with_fans,
        "partial": bool(warn_list) or detail_fail > 0,
        "warnings": warn_list,
        "cost": cost,
        "note": note,
    }


def main() -> None:
    raw = sys.argv[1] if len(sys.argv) > 1 else ""
    try:
        params = json.loads(raw) if str(raw).strip() else {}
        if not isinstance(params, dict):
            raise LinkDetailError("参数必须是 JSON 对象")
        data = run(params)
        print(json.dumps({"ok": True, "data": data}, ensure_ascii=False))
    except XhsError as err:
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": str(err),
                    "hint": "提供 input_path；确认 SKILL_TIKHUB_API_KEY / SKILL_TIKHUB_BASE_URL；59 需 SKILL_TIKHUB_PROXY；402 表示该路由要付费余额",
                },
                ensure_ascii=False,
            )
        )
    except Exception as err:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": f"{type(err).__name__}: {err}"}, ensure_ascii=False))


if __name__ == "__main__":
    main()
